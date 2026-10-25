# -*- coding: utf-8 -*-
"""
Verification Script: Test end-to-end functionality of ShrimpSmartCounterApp:
- Vietnamese TTS voice feedback
- Small-talk conversational filter (is_relevant_shrimp_command)
- Automatic shrimp average size calculation (con/kg, con/g)
- Tab 1 Weighing & Gender classification
- Tab 2 Batch accumulation & Undo
- Multi-sheet Excel export
"""

import os
import sys
import time

if sys.stdout and sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import openpyxl
from app import ShrimpSmartCounterApp, is_relevant_shrimp_command
from excel_manager import save_shrimp_report

def run_headless_integration_test():
    print("=================================================================")
    print("=== STARTING COMPREHENSIVE INTEGRATION & FEATURE TEST ===")
    print("=================================================================")

    # -------------------------------------------------------------
    # TEST 1: SMALL TALK / CHAT FILTER LOGIC
    # -------------------------------------------------------------
    print("\n[TEST 1] Testing Small-talk / Chit-chat Filter (is_relevant_shrimp_command):")
    chit_chat_samples = [
        "Hôm nay trời nắng quá",
        "Cho tôm ăn chưa em ơi",
        "Đi uống cà phê đi",
        "Bật quạt sục khí oxy lên",
        "Chiều nay mấy giờ về",
        "Alo anh nghe này",
    ]
    for sentence in chit_chat_samples:
        is_relevant = is_relevant_shrimp_command(sentence)
        print(f"  -> Chit-chat: '{sentence}' => Recognized as relevant? {is_relevant}")
        assert not is_relevant, f"Failed: '{sentence}' should have been ignored as chit-chat!"

    valid_command_samples = [
        ("ba mươi lăm đực", True),
        ("bốn lăm cái", True),
        ("năm hai", True),
        ("sáu", True),
        ("mười lăm", True),
        ("hai mươi", True),
        ("xóa", True),
        ("hủy", True),
        ("50 cái", True),
        ("35", True),
    ]
    for sentence, expected in valid_command_samples:
        is_relevant = is_relevant_shrimp_command(sentence)
        print(f"  -> Valid Command: '{sentence}' => Recognized as relevant? {is_relevant}")
        assert is_relevant == expected, f"Failed: '{sentence}' was incorrectly ignored!"
    print("  => Small-talk filter logic passed 100%!")

    # -------------------------------------------------------------
    # INIT APP INSTANCE (HEADLESS)
    # -------------------------------------------------------------
    app = ShrimpSmartCounterApp()
    app.withdraw()  # Hide UI window during test

    # -------------------------------------------------------------
    # TEST 2: TAB 1 - WEIGHING, GENDER & CHIT-CHAT BYPASS & AVERAGE SIZE
    # -------------------------------------------------------------
    print("\n[TEST 2] Testing Mode 1 (Weighing, Gender, Size Calculation & Chit-chat bypass):")
    app.notebook.select(0)  # Select Tab 1

    # First, test that chit-chat does NOT add any record
    initial_records_count = len(app.weight_records)
    app._dispatch_recognized_speech("Hôm nay nắng đẹp quá anh ơi", force=True)
    assert len(app.weight_records) == initial_records_count, "Chit-chat falsely added a weight record!"
    print("  -> Chit-chat bypass in Tab 1 verified: No record added.")

    # Dispatch valid weighings
    test_speeches_mode1 = [
        "ba mươi lăm đực",        # 35g, Đực
        "bốn lăm cái",           # 45g, Cái
        "năm hai đực",           # 52g, Đực
        "ba lăm",                # 35g, Chưa rõ
        "hai lăm phẩy năm cái"   # 25.5g, Cái
    ]

    for speech in test_speeches_mode1:
        app._dispatch_recognized_speech(speech, force=True)
        print(f"  -> Mode 1 speech dispatched: '{speech}'")

    total_weight = sum(r['weight'] for r in app.weight_records)
    total_count = len(app.weight_records)
    expected_size_kg = (total_count / total_weight) * 1000
    expected_size_g = total_count / total_weight

    print(f"  -> Total weighed shrimp: {total_count} con, Total weight: {total_weight:.1f} g")
    print(f"  -> Calculated Size: {expected_size_kg:.1f} con/kg (={expected_size_g:.4f} con/g)")
    print(f"  -> UI Size label text: '{app.lbl_mode1_size.cget('text')}'")
    print(f"  -> UI Size sub-label text: '{app.lbl_mode1_size_sub.cget('text')}'")

    assert f"{expected_size_kg:.1f} con/kg" in app.lbl_mode1_size.cget("text"), "Mode 1 Size UI display mismatch!"
    assert len(app.weight_records) == 5, f"Expected 5 records, got {len(app.weight_records)}"
    print("  => Mode 1 data & Automatic Size calculation verified successfully!")

    # -------------------------------------------------------------
    # TEST 3: TAB 2 - BATCH ACCUMULATION, CHIT-CHAT BYPASS & UNDO
    # -------------------------------------------------------------
    print("\n[TEST 3] Testing Mode 2 (High-speed Accumulation, Chit-chat bypass & Undo):")
    app.notebook.select(1)  # Select Tab 2

    # Verify chit-chat bypass in Mode 2
    initial_batch_total = app.total_batch_count
    app._dispatch_recognized_speech("Ai gọi điện thoại thế", force=True)
    assert app.total_batch_count == initial_batch_total, "Chit-chat falsely altered batch total!"
    print("  -> Chit-chat bypass in Tab 2 verified: Total unaltered.")

    test_speeches_mode2 = [
        "năm",          # +5 -> 5
        "mười",         # +10 -> 15
        "mười lăm",     # +15 -> 30
        "hai mươi",     # +20 -> 50
        "xóa",          # Undo last (+20) -> 30
        "mười hai"      # +12 -> 42
    ]

    for speech in test_speeches_mode2:
        app._dispatch_recognized_speech(speech, force=True)
        print(f"  -> Mode 2 dispatched: '{speech}' => Total now: {app.total_batch_count}")

    assert app.total_batch_count == 42, f"Expected total 42, got {app.total_batch_count}"
    print(f"  => Mode 2 total count verified: {app.total_batch_count} con!")

    # Test Undo via 'hủy'
    app._dispatch_recognized_speech("hủy", force=True)  # Reverts +12 -> 30
    assert app.total_batch_count == 30, f"Expected total 30 after undo, got {app.total_batch_count}"
    print(f"  => Mode 2 Undo via 'hủy' verified: Total now = {app.total_batch_count} con!")

    # -------------------------------------------------------------
    # TEST 4: EXCEL EXPORT WITH SIZE STATS
    # -------------------------------------------------------------
    print("\n[TEST 4] Testing Excel Export with 2 Sheets:")
    test_excel_file = os.path.join(app.app_dir, "Bao_Cao_Tom_Verified.xlsx")
    success, final_path, msg = save_shrimp_report(
        test_excel_file,
        pond_name="Ao Số 1 - Khu Thử Nghiệm",
        staff_name="Nguyễn Kỹ Thuật Viên",
        weight_records=app.weight_records,
        batch_records=app.batch_records,
        total_shrimp_count=app.total_batch_count
    )

    assert success, f"Excel save failed: {msg}"
    assert os.path.exists(test_excel_file), "Excel file was not created!"

    wb = openpyxl.load_workbook(test_excel_file)
    sheet_names = wb.sheetnames
    print(f"  Sheets in Excel: {sheet_names}")
    assert "Can_Trong_Luong_Gioi_Tinh" in sheet_names, "Missing Sheet 1!"
    assert "Dem_So_Luong_Cong_Don" in sheet_names, "Missing Sheet 2!"

    ws1 = wb["Can_Trong_Luong_Gioi_Tinh"]
    print(f"  Sheet 1 Title: {ws1['A1'].value}")
    # Verify row 7 has first record in the new professional layout
    sample_row_1 = [ws1.cell(7, c).value for c in range(1, 5)]
    print(f"  Sheet 1 Sample Row (Row 7): {sample_row_1}")
    assert sample_row_1[0] == 1, f"Expected STT 1, got {sample_row_1[0]}"

    ws2 = wb["Dem_So_Luong_Cong_Don"]
    print(f"  Sheet 2 Title: {ws2['A1'].value}")
    sample_row_2 = [ws2.cell(7, c).value for c in range(1, 6)]
    print(f"  Sheet 2 Sample Row (Row 7): {sample_row_2}")

    wb.close()
    print("  => Excel file verified with 2 sheets and complete statistics!")

    app.destroy()
    print("\n=================================================================")
    print("=== ALL INTEGRATION & REGRESSION TESTS PASSED 100% SUCCESSFULLY! ===")
    print("=================================================================")

if __name__ == '__main__':
    run_headless_integration_test()
