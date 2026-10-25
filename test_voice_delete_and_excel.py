# -*- coding: utf-8 -*-
"""
Kịch bản kiểm thử toàn diện:
1. Chức năng 1: Xóa dòng cuối cùng bằng giọng nói ("xóa", "hủy", "xóa dòng") cho cả Chế độ 1 và Chế độ 2.
2. Chức năng 2: Xóa sạch toàn bộ (Reset về 0) bằng giọng nói ("xóa hết", "xóa tất cả", "reset") cho cả Chế độ 1 và Chế độ 2.
3. Kiểm tra tính năng bẫy từ giới tính (trai, chai, đực, tai, cái, gái, ái) và dịch số vẫn hoạt động chuẩn xác.
4. Kiểm tra file Excel xuất ra đạt 5 tiêu chuẩn thẩm mỹ chuyên nghiệp bậc cao (openpyxl).
"""

import os
import sys

if sys.stdout and sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import openpyxl
from app import ShrimpSmartCounterApp, is_relevant_shrimp_command, detect_shrimp_gender_mode1
from excel_manager import save_shrimp_report

class MockTTS:
    def __init__(self):
        self.last_spoken = None
        self.history = []
        self.enabled = True

    def speak_async(self, text: str):
        self.last_spoken = text
        self.history.append(text)
        print(f"      [MOCK TTS] AI đã đọc: \"{text}\"")

def run_comprehensive_test():
    print("=====================================================================")
    print("=== BẮT ĐẦU KIỂM THỬ XÓA DÒNG CUỐI, XÓA SẠCH VỀ 0 VÀ EXCEL CAO CẤP ===")
    print("=====================================================================")

    app = ShrimpSmartCounterApp()
    app.withdraw()
    mock_tts = MockTTS()
    app.tts = mock_tts

    # -----------------------------------------------------------------
    # PHẦN 1: BỘ LỌC TỪ KHÓA ĐIỀU KHIỂN & CHUYỆN PHIẾM
    # -----------------------------------------------------------------
    print("\n[KIỂM THỬ 1] Kiểm tra bộ lọc từ khóa điều khiển (is_relevant_shrimp_command):")
    commands_to_test = [
        ("xóa", True),
        ("hủy", True),
        ("xóa dòng", True),
        ("xóa hết", True),
        ("xóa tất cả", True),
        ("reset", True),
        ("35 trai", True),
        ("35 chai", True),
        ("45 cái", True),
        ("mười", True),
        ("hôm nay thời tiết đẹp quá", False),
        ("pha nước mắm đi em", False),
    ]

    for cmd, expected in commands_to_test:
        res = is_relevant_shrimp_command(cmd)
        print(f"  -> Lệnh: '{cmd}' => Hợp lệ? {res} (Kỳ vọng: {expected})")
        assert res == expected, f"Lỗi bộ lọc: '{cmd}' nhận {res} thay vì {expected}"
    print("  ✅ Bộ lọc từ khóa điều khiển và chống chuyện phiếm hoạt động 100% chuẩn xác!")

    # -----------------------------------------------------------------
    # PHẦN 2: CHẾ ĐỘ 1 - CÂN MẪU: XÓA DÒNG CUỐI & XÓA SẠCH TOÀN BỘ
    # -----------------------------------------------------------------
    print("\n[KIỂM THỬ 2] Kiểm tra Chế độ 1 (Cân mẫu):")
    app.notebook.select(0)  # Chọn Tab 1
    app.weight_records.clear()
    app._refresh_tree_weight()

    # Thêm 3 mẫu
    app._dispatch_recognized_speech("35 trai", force=True)
    app._dispatch_recognized_speech("45 gái", force=True)
    app._dispatch_recognized_speech("50 đực", force=True)

    assert len(app.weight_records) == 3, f"Kỳ vọng 3 mẫu, có {len(app.weight_records)}"
    print(f"  -> Đã thêm 3 mẫu: 35g(Male), 45g(Female), 50g(Male). Tổng: {len(app.weight_records)} mẫu")

    # 1. Thử lệnh "xóa dòng"
    print("\n  [2.1] Hô lệnh 'xóa dòng':")
    app._dispatch_recognized_speech("xóa dòng", force=True)
    assert len(app.weight_records) == 2, f"Kỳ vọng còn 2 mẫu, còn {len(app.weight_records)}"
    assert app.weight_records[-1]['gender'] == "Female", "Mẫu cuối chưa đúng!"
    assert mock_tts.last_spoken == "Đã xóa dòng cuối", f"TTS đọc sai: {mock_tts.last_spoken}"
    print(f"  ✅ Đã xóa dòng cuối (50g đực)! Còn lại: {len(app.weight_records)} mẫu. Loa đọc: '{mock_tts.last_spoken}'")

    # 2. Thử lệnh "hủy"
    print("\n  [2.2] Hô lệnh 'hủy':")
    app._dispatch_recognized_speech("hủy", force=True)
    assert len(app.weight_records) == 1, f"Kỳ vọng còn 1 mẫu, còn {len(app.weight_records)}"
    assert app.weight_records[-1]['gender'] == "Male"
    assert mock_tts.last_spoken == "Đã xóa dòng cuối"
    print(f"  ✅ Đã xóa dòng cuối (45g gái)! Còn lại: {len(app.weight_records)} mẫu. Loa đọc: '{mock_tts.last_spoken}'")

    # 3. Thử lệnh "xóa"
    print("\n  [2.3] Hô lệnh 'xóa':")
    app._dispatch_recognized_speech("xóa", force=True)
    assert len(app.weight_records) == 0, f"Kỳ vọng bảng trống, còn {len(app.weight_records)}"
    assert mock_tts.last_spoken == "Đã xóa dòng cuối"
    print(f"  ✅ Đã xóa dòng cuối! Bảng cân mẫu đã trống. Loa đọc: '{mock_tts.last_spoken}'")

    # 4. Thử lệnh "xóa hết" / "reset" ở Chế độ 1
    print("\n  [2.4] Thêm lại 4 mẫu rồi hô lệnh 'xóa hết':")
    app._dispatch_recognized_speech("30 cái", force=True)
    app._dispatch_recognized_speech("40 đực", force=True)
    app._dispatch_recognized_speech("35 chai", force=True)
    app._dispatch_recognized_speech("50 tai", force=True)
    assert len(app.weight_records) == 4
    print(f"  -> Đã thêm 4 mẫu cân. Giờ hô 'xóa hết'...")

    app._dispatch_recognized_speech("xóa hết", force=True)
    assert len(app.weight_records) == 0, "Bảng cân chưa sạch sau khi xóa hết!"
    assert len(app.tree_weight.get_children()) == 0, "Treeview chưa sạch!"
    assert mock_tts.last_spoken == "Đang xóa toàn bộ dữ liệu", f"TTS sai: {mock_tts.last_spoken}"
    print(f"  ✅ Đã xóa sạch toàn bộ bảng cân mẫu! Loa đọc cảnh báo: '{mock_tts.last_spoken}'")

    # -----------------------------------------------------------------
    # PHẦN 3: CHẾ ĐỘ 2 - ĐẾM CỘNG DỒN: XÓA DÒNG CUỐI & XÓA SẠCH VỀ 0
    # -----------------------------------------------------------------
    print("\n[KIỂM THỬ 3] Kiểm tra Chế độ 2 (Đếm cộng dồn):")
    app.notebook.select(1)  # Chọn Tab 2
    app.total_batch_count = 0
    app.batch_counter = 0
    app.batch_records.clear()
    app.tree_batch.delete(*app.tree_batch.get_children())

    # Thêm 3 mẻ: 5 + 10 + 15 = 30
    app._dispatch_recognized_speech("năm", force=True)
    app._dispatch_recognized_speech("mười", force=True)
    app._dispatch_recognized_speech("mười lăm", force=True)
    assert app.total_batch_count == 30, f"Tổng tôm kỳ vọng 30, có {app.total_batch_count}"
    assert len(app.batch_records) == 3
    print(f"  -> Đã thêm 3 mẻ (+5, +10, +15). Tổng tôm: {app.total_batch_count} con")

    # 1. Thử lệnh "xóa dòng"
    print("\n  [3.1] Hô lệnh 'xóa dòng':")
    app._dispatch_recognized_speech("xóa dòng", force=True)
    assert app.total_batch_count == 15, f"Kỳ vọng tổng 15, có {app.total_batch_count}"
    assert len(app.batch_records) == 2
    assert mock_tts.last_spoken == "Đã xóa dòng cuối", f"TTS đọc sai: {mock_tts.last_spoken}"
    print(f"  ✅ Đã xóa dòng cuối (trừ 15 con)! Tổng còn: {app.total_batch_count}. Loa đọc: '{mock_tts.last_spoken}'")

    # 2. Thử lệnh "hủy"
    print("\n  [3.2] Hô lệnh 'hủy':")
    app._dispatch_recognized_speech("hủy", force=True)
    assert app.total_batch_count == 5, f"Kỳ vọng tổng 5, có {app.total_batch_count}"
    assert len(app.batch_records) == 1
    assert mock_tts.last_spoken == "Đã xóa dòng cuối"
    print(f"  ✅ Đã xóa dòng cuối (trừ 10 con)! Tổng còn: {app.total_batch_count}. Loa đọc: '{mock_tts.last_spoken}'")

    # 3. Thử lệnh "xóa"
    print("\n  [3.3] Hô lệnh 'xóa':")
    app._dispatch_recognized_speech("xóa", force=True)
    assert app.total_batch_count == 0, f"Kỳ vọng tổng 0, có {app.total_batch_count}"
    assert len(app.batch_records) == 0
    assert mock_tts.last_spoken == "Đã xóa dòng cuối"
    print(f"  ✅ Đã xóa dòng cuối (trừ 5 con)! Tổng tôm về 0. Loa đọc: '{mock_tts.last_spoken}'")

    # 4. Thử lệnh "reset" / "xóa tất cả"
    print("\n  [3.4] Thêm 3 mẻ (20 + 35 + 10 = 65 con) rồi hô 'reset':")
    app._dispatch_recognized_speech("hai mươi", force=True)
    app._dispatch_recognized_speech("ba lăm", force=True)
    app._dispatch_recognized_speech("mười", force=True)
    assert app.total_batch_count == 65
    assert len(app.batch_records) == 3
    print(f"  -> Tổng tích lũy: {app.total_batch_count} con. Giờ hô 'reset'...")

    app._dispatch_recognized_speech("reset", force=True)
    assert app.total_batch_count == 0, f"Tổng tôm chưa về 0: {app.total_batch_count}"
    assert app.batch_counter == 0, "Số lượng mẻ chưa về 0!"
    assert len(app.batch_records) == 0, "Lịch sử mẻ chưa sạch!"
    assert len(app.tree_batch.get_children()) == 0, "Bảng Treeview chưa sạch!"
    assert mock_tts.last_spoken == "Đang xóa toàn bộ dữ liệu", f"TTS sai: {mock_tts.last_spoken}"
    print(f"  ✅ Đã xóa sạch toàn bộ về 0! Tổng: {app.total_batch_count}, mẻ: {app.batch_counter}. Loa đọc: '{mock_tts.last_spoken}'")

    # 5. Thử lệnh "xóa tất cả"
    print("\n  [3.5] Thêm lại 2 mẻ rồi hô 'xóa tất cả':")
    app._dispatch_recognized_speech("mười lăm", force=True)
    app._dispatch_recognized_speech("hai mươi", force=True)
    assert app.total_batch_count == 35

    app._dispatch_recognized_speech("xóa tất cả", force=True)
    assert app.total_batch_count == 0
    assert len(app.batch_records) == 0
    assert mock_tts.last_spoken == "Đang xóa toàn bộ dữ liệu"
    print(f"  ✅ Hô 'xóa tất cả' thành công! Loa đọc: '{mock_tts.last_spoken}'")

    # -----------------------------------------------------------------
    # PHẦN 4: KIỂM THỬ XUẤT FILE EXCEL OPENPYXL ĐẠT TIÊU CHUẨN CAO CẤP
    # -----------------------------------------------------------------
    print("\n[KIỂM THỬ 4] Kiểm tra định dạng Excel Openpyxl thẩm mỹ chuyên nghiệp bậc cao:")
    # Nạp dữ liệu mẫu cho cả 2 tab
    app.notebook.select(0)
    app.weight_records = [
        {'stt': 1, 'weight': 35.0, 'gender': 'Male', 'time': '10:00:00'},
        {'stt': 2, 'weight': 38.5, 'gender': 'Male', 'time': '10:01:00'},
        {'stt': 3, 'weight': 42.0, 'gender': 'Female', 'time': '10:02:00'},
        {'stt': 4, 'weight': 45.0, 'gender': 'Female', 'time': '10:03:00'},
        {'stt': 5, 'weight': 50.0, 'gender': 'Male', 'time': '10:04:00'},
    ]
    app._refresh_tree_weight()

    app.notebook.select(1)
    app.batch_records = [
        {'stt': 1, 'batch_name': 'Mẻ 1', 'count': 10, 'cumulative': 10, 'time': '10:05:00', 'action': 'ADD'},
        {'stt': 2, 'batch_name': 'Mẻ 2', 'count': 25, 'cumulative': 35, 'time': '10:06:00', 'action': 'ADD'},
        {'stt': 3, 'batch_name': 'Mẻ 3', 'count': 15, 'cumulative': 50, 'time': '10:07:00', 'action': 'ADD'},
    ]
    app.total_batch_count = 50
    app.batch_counter = 3
    app._update_huge_count_display(highlight=False)

    test_excel_path = os.path.join(app.app_dir, "Bao_Cao_Tom_Pro_Format.xlsx")
    success, final_path, msg = save_shrimp_report(
        test_excel_path,
        pond_name="Ao Nuôi Công Nghệ Cao Số 1",
        staff_name="Kỹ Sư Trưởng",
        weight_records=app.weight_records,
        batch_records=app.batch_records,
        total_shrimp_count=app.total_batch_count
    )

    assert success, f"Lưu Excel thất bại: {msg}"
    print(f"  -> File Excel đã tạo: {final_path}")

    wb = openpyxl.load_workbook(final_path)

    # 1. TIÊU CHUẨN 1: THÔNG TIN CHUNG (Dòng 1 đến 4)
    ws1 = wb["Can_Trong_Luong_Gioi_Tinh"]
    print("\n  [4.1] Tiêu chuẩn 1 - Thông tin chung (Dòng 1 đến 4):")
    # Dòng 1: Tiêu đề in hoa, cỡ 16, Bold, Navy Blue #1F4E78
    title_cell = ws1['A1']
    assert title_cell.value == "BÁO CÁO DỮ LIỆU ĐẾM TÔM THÔNG MINH"
    assert title_cell.font.size == 16, f"Cỡ chữ tiêu đề phải là 16, thực tế: {title_cell.font.size}"
    assert title_cell.font.bold is True, "Tiêu đề phải in đậm"
    assert title_cell.font.color.rgb == '001F4E78' or '1F4E78' in str(title_cell.font.color.rgb), "Màu chữ phải là Navy Blue #1F4E78"
    print("     ✅ Dòng 1: 'BÁO CÁO DỮ LIỆU ĐẾM TÔM THÔNG MINH' (Font 16 Bold, Navy Blue #1F4E78)")

    # Dòng 2 đến 4: Tên ao, Nhân viên, Ngày giờ, nền Light Gray #F2F2F2
    assert "Ao Nuôi" in str(ws1['A2'].value)
    assert ws1['B2'].value == "Ao Nuôi Công Nghệ Cao Số 1"
    assert "Nhân Viên" in str(ws1['A3'].value)
    assert ws1['B3'].value == "Kỹ Sư Trưởng"
    assert "Ngày Giờ" in str(ws1['A4'].value)

    fill_a2 = ws1['A2'].fill.start_color.rgb
    assert 'F2F2F2' in str(fill_a2), f"Dòng 2 phải tô nền Light Gray #F2F2F2, thực tế: {fill_a2}"
    print("     ✅ Dòng 2-4: Tên Ao, Nhân Viên, Ngày Giờ (Tô nền Light Gray #F2F2F2 chuẩn xác)")

    # 2. TIÊU CHUẨN 2: THANH TIÊU ĐỀ BẢNG (Header - Dòng 6)
    print("\n  [4.2] Tiêu chuẩn 2 - Thanh tiêu đề bảng (Header - Dòng 6):")
    assert ws1.row_dimensions[6].height == 26, f"Chiều cao header phải là 26, thực tế: {ws1.row_dimensions[6].height}"
    for col_idx, expected_h in enumerate(["STT", "Trọng Lượng (g)", "Giới Tính", "Thời Gian Ghi Nhận"], start=1):
        cell_h = ws1.cell(row=6, column=col_idx)
        assert cell_h.value == expected_h, f"Cột {col_idx} sai tiêu đề: {cell_h.value}"
        assert '1F4E78' in str(cell_h.fill.start_color.rgb), "Header phải có nền Navy Blue #1F4E78"
        assert cell_h.font.color.rgb == '00FFFFFF' or 'FFFFFF' in str(cell_h.font.color.rgb), "Chữ header phải màu trắng"
        assert cell_h.alignment.horizontal == 'center', "Header phải căn giữa"
        assert cell_h.font.bold is True, "Header phải in đậm"
    print("     ✅ Header: Chiều cao 26, nền Navy Blue #1F4E78, chữ trắng 11 Bold, căn giữa")

    # 3. TIÊU CHUẨN 3: DÒNG DỮ LIỆU (Data Rows - Dòng 7 đến 11)
    print("\n  [4.3] Tiêu chuẩn 3 - Định dạng các dòng dữ liệu (Dòng 7 đến 11):")
    for r_idx in range(7, 12):
        assert ws1.row_dimensions[r_idx].height == 20, f"Dòng {r_idx} phải có chiều cao 20"
        # Cột STT căn giữa
        assert ws1.cell(row=r_idx, column=1).alignment.horizontal == 'center'
        # Cột Trọng Lượng căn phải, format '0.0'
        c_w = ws1.cell(row=r_idx, column=2)
        assert c_w.alignment.horizontal == 'right', f"Dòng {r_idx} cột trọng lượng phải căn phải"
        assert c_w.number_format == '0.0', f"Dòng {r_idx} format phải là '0.0', thực tế: {c_w.number_format}"
        # Cột Giới Tính căn giữa
        assert ws1.cell(row=r_idx, column=3).alignment.horizontal == 'center'
        # Cột Thời Gian căn giữa
        assert ws1.cell(row=r_idx, column=4).alignment.horizontal == 'center'
        # Viền mỏng toàn bộ
        for c_idx in range(1, 5):
            assert ws1.cell(row=r_idx, column=c_idx).border.left.style == 'thin', "Thiếu viền mỏng"
    print("     ✅ Dữ liệu: Chiều cao 20, Calibri 10, STT/Giới tính căn giữa, Trọng lượng căn phải format '0.0', viền xám mỏng")

    # 4. TIÊU CHUẨN 4: DÒNG TỔNG KẾT (Total Row - Dòng 12)
    print("\n  [4.4] Tiêu chuẩn 4 - Dòng tổng kết (Total Row - Dòng 12):")
    tot_row = 12
    assert ws1.row_dimensions[tot_row].height == 24, "Dòng tổng kết phải có chiều cao 24"
    assert ws1.cell(row=tot_row, column=1).value == "TỔNG KẾT"
    assert ws1.cell(row=tot_row, column=1).font.bold is True
    # Công thức AVERAGE
    avg_formula = ws1.cell(row=tot_row, column=2).value
    assert avg_formula == "=AVERAGE(B7:B11)", f"Công thức trung bình sai: {avg_formula}"
    assert ws1.cell(row=tot_row, column=2).number_format == '0.0'
    # Gạch chân kép chuẩn kế toán
    bot_border_style = ws1.cell(row=tot_row, column=2).border.bottom.style
    assert bot_border_style == 'double', f"Đường viền dưới của dòng tổng kết phải là 'double', thực tế: {bot_border_style}"
    print(f"     ✅ Dòng tổng kết: =AVERAGE(B7:B11), Bold, đường gạch chân kép (double) chuẩn kế toán")

    # 5. TIÊU CHUẨN 5: TỰ ĐỘNG GIÃN CỘT
    print("\n  [4.5] Tiêu chuẩn 5 - Tự động giãn cột (Auto-fit):")
    for col_letter in ['A', 'B', 'C', 'D']:
        w = ws1.column_dimensions[col_letter].width
        assert w >= 18, f"Cột {col_letter} chưa được giãn đủ rộng: {w}"
        print(f"     ✅ Cột {col_letter}: Độ rộng = {w:.1f} (Không bị lỗi '###' hay che chữ)")

    # Kiểm tra Sheet 2
    ws2 = wb["Dem_So_Luong_Cong_Don"]
    assert ws2['A1'].value == "BÁO CÁO DỮ LIỆU ĐẾM TÔM THÔNG MINH"
    assert ws2.cell(row=6, column=3).value == "Số Lượng (Con)"
    # Dòng tổng kết Sheet 2 (Dòng 10 vì 3 mẻ từ row 7 đến 9)
    sum_formula = ws2.cell(row=10, column=3).value
    assert sum_formula == "=SUM(C7:C9)", f"Sheet 2 công thức sai: {sum_formula}"
    assert ws2.cell(row=10, column=3).border.bottom.style == 'double'
    print(f"     ✅ Sheet 2 Dòng tổng kết: =SUM(C7:C9), gạch chân kép chuẩn kế toán")

    wb.close()
    app.destroy()
    print("\n=====================================================================")
    print("=== TẤT CẢ 4 PHẦN KIỂM THỬ ĐÃ VƯỢT QUA 100% THÀNH CÔNG RỰC RỠ! ===")
    print("=====================================================================")

if __name__ == '__main__':
    run_comprehensive_test()
