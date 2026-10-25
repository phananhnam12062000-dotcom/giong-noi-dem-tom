# -*- coding: utf-8 -*-
"""
Test script to verify the new gender keyword requirements in Mode 1:
1. "trai", "đực", "tai" -> "Male"
2. "cái", "gái", "ái" -> "Female"
3. Neither -> "Chưa rõ"
4. Display and save "Male" and "Female" into UI table and Excel.
"""

import os
import sys

if sys.stdout and sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import openpyxl
from app import ShrimpSmartCounterApp, detect_shrimp_gender_mode1
from vietnamese_parser import parse_mode1_weighing_and_gender
from excel_manager import save_shrimp_report

def test_new_gender_mapping():
    print("=====================================================================")
    print("=== BẮT ĐẦU KIỂM THỬ TỪ KHÓA GIỚI TÍNH MỚI (TRAI/ĐỰC/TAI & CÁI/GÁI/ÁI) ===")
    print("=====================================================================")

    # 1. Kiểm tra Giống Đực: "trai", "chai", "đực", "tai" -> "Male"
    male_cases = [
        ("35 trai", 35.0, "Male"),
        ("35 chai", 35.0, "Male"),
        ("40 đực", 40.0, "Male"),
        ("50 tai", 50.0, "Male"),
        ("ba lăm trai", 35.0, "Male"),
        ("ba lăm chai", 35.0, "Male"),
        ("bốn lăm đực", 45.0, "Male"),
        ("năm hai tai", 52.0, "Male"),
        ("35trai", 35.0, "Male"),
        ("35chai", 35.0, "Male"),
        ("50tai", 50.0, "Male"),
        ("40đực", 40.0, "Male"),
    ]

    print("\n[TEST 1] Kiểm tra từ khóa Giống ĐỰC ('trai', 'chai', 'đực', 'tai' -> 'Male'):")
    for text, exp_weight, exp_gender in male_cases:
        gender = detect_shrimp_gender_mode1(text)
        print(f"  -> '{text}' => Giới tính: '{gender}' (Kỳ vọng: '{exp_gender}')")
        assert gender == exp_gender, f"Lỗi: '{text}' trả về '{gender}', mong đợi '{exp_gender}'"

    # 2. Kiểm tra Giống Cái: "cái", "gái", "ái" -> "Female"
    female_cases = [
        ("45 cái", 45.0, "Female"),
        ("38 gái", 38.0, "Female"),
        ("25 ái", 25.0, "Female"),
        ("ba lăm cái", 35.0, "Female"),
        ("bốn lăm gái", 45.0, "Female"),
        ("hai lăm ái", 25.0, "Female"),
        ("45cái", 45.0, "Female"),
        ("38gái", 38.0, "Female"),
        ("25ái", 25.0, "Female"),
    ]

    print("\n[TEST 2] Kiểm tra từ khóa Giống CÁI ('cái', 'gái', 'ái' -> 'Female'):")
    for text, exp_weight, exp_gender in female_cases:
        gender = detect_shrimp_gender_mode1(text)
        print(f"  -> '{text}' => Giới tính: '{gender}' (Kỳ vọng: '{exp_gender}')")
        assert gender == exp_gender, f"Lỗi: '{text}' trả về '{gender}', mong đợi '{exp_gender}'"

    # 3. Kiểm tra không có từ khóa -> "Chưa rõ"
    unspecified_cases = [
        ("35", 35.0, "Chưa rõ"),
        ("ba mươi lăm", 35.0, "Chưa rõ"),
        ("50 gram", 50.0, "Chưa rõ"),
        ("hai mươi lăm phẩy năm", 25.5, "Chưa rõ")
    ]

    print("\n[TEST 3] Kiểm tra khi không đọc từ khóa giới tính -> Giữ nguyên 'Chưa rõ':")
    for text, exp_weight, exp_gender in unspecified_cases:
        gender = detect_shrimp_gender_mode1(text)
        print(f"  -> '{text}' => Giới tính: '{gender}' (Kỳ vọng: '{exp_gender}')")
        assert gender == exp_gender, f"Lỗi: '{text}' trả về '{gender}', mong đợi '{exp_gender}'"

    # 4. Kiểm thử tích hợp UI Treeview và Xuất Excel
    print("\n[TEST 4] Kiểm thử tích hợp lưu 'Male', 'Female' vào Treeview và Báo cáo Excel:")
    app = ShrimpSmartCounterApp()
    app.withdraw()
    app.notebook.select(0)  # Tab 1
    app.tts.enabled = False  # Tắt loa để test nhanh

    test_samples = [
        ("35 trai", 35.0, "Male"),
        ("38 chai", 38.0, "Male"),
        ("45 gái", 45.0, "Female"),
        ("40 đực", 40.0, "Male"),
        ("30 cái", 30.0, "Female"),
        ("50 tai", 50.0, "Male"),
        ("28 ái", 28.0, "Female"),
        ("55", 55.0, "Chưa rõ"),
    ]

    for speech, exp_weight, exp_gender in test_samples:
        app._handle_mode1_speech(speech)
        last_rec = app.weight_records[-1]
        print(f"  -> Đọc: '{speech}' => Đã lưu: {last_rec['weight']}g - '{last_rec['gender']}'")
        assert last_rec['weight'] == exp_weight, f"Sai trọng lượng: {last_rec['weight']} != {exp_weight}"
        assert last_rec['gender'] == exp_gender, f"Sai giới tính: {last_rec['gender']} != {exp_gender}"

    # Kiểm tra các dòng trong bảng Treeview
    tree_rows = app.tree_weight.get_children()
    assert len(tree_rows) == len(test_samples), f"Bảng Treeview thiếu dòng: {len(tree_rows)}"

    for idx, item in enumerate(tree_rows):
        values = app.tree_weight.item(item, "values")
        stt, w_str, g_val, t_val = values
        expected_g = test_samples[idx][2]
        assert g_val == expected_g, f"Dòng {idx} hiển thị '{g_val}' thay vì '{expected_g}'"
        print(f"     ✅ Bảng Treeview dòng {stt}: {w_str}g | Giới Tính = '{g_val}'")

    # Kiểm tra nhãn tỷ lệ trên giao diện
    ratio_text = app.lbl_mode1_ratio.cget("text")
    print(f"  -> Nhãn thống kê tỷ lệ trên UI: '{ratio_text}'")
    assert "Male: 4" in ratio_text, "Thống kê Male trên UI sai lệch!"
    assert "Female: 3" in ratio_text, "Thống kê Female trên UI sai lệch!"

    # 5. Kiểm tra file Excel xuất ra có chứa cột Giới Tính là 'Male' và 'Female'
    print("\n[TEST 5] Kiểm tra file Excel xuất ra:")
    test_excel_path = os.path.join(app.app_dir, "Bao_Cao_Tom_Verified.xlsx")
    success, final_path, msg = save_shrimp_report(
        test_excel_path,
        pond_name="Ao Thử Nghiệm Giới Tính",
        staff_name="Kỹ Thuật Viên",
        weight_records=app.weight_records,
        batch_records=app.batch_records,
        total_shrimp_count=app.total_batch_count
    )
    assert success, f"Lưu Excel thất bại: {msg}"

    wb = openpyxl.load_workbook(test_excel_path)
    ws1 = wb["Can_Trong_Luong_Gioi_Tinh"]

    # Kiểm tra dòng tiêu đề đạt chuẩn thẩm mỹ mới
    assert ws1['A1'].value == "BÁO CÁO DỮ LIỆU ĐẾM TÔM THÔNG MINH", "Tiêu đề Excel sai lệch!"
    assert ws1['A2'].value == "Tên Ao Nuôi / Mã Lô:", "Dòng 2 thiếu thông tin tên ao!"
    assert ws1['A3'].value == "Nhân Viên Cân:", "Dòng 3 thiếu thông tin nhân viên!"
    assert ws1['A4'].value == "Ngày Giờ Xuất Báo Cáo:", "Dòng 4 thiếu ngày giờ!"
    assert ws1['A6'].value == "STT", "Dòng 6 phải là Header STT!"

    # Kiểm tra dòng dữ liệu từ row 7 đến row 14
    for row_idx, exp_data in enumerate(test_samples, start=7):
        cell_gender = ws1.cell(row=row_idx, column=3).value
        assert cell_gender == exp_data[2], f"Excel row {row_idx} có giới tính '{cell_gender}' thay vì '{exp_data[2]}'"
        print(f"     ✅ Excel Dòng {row_idx}: Cột Giới Tính = '{cell_gender}'")

    # Kiểm tra dòng tổng kết (Total Row) có công thức Excel và gạch chân kép
    total_row_idx = 7 + len(test_samples)
    tot_title = ws1.cell(row=total_row_idx, column=1).value
    tot_formula = ws1.cell(row=total_row_idx, column=2).value
    print(f"  -> Dòng tổng kết Row {total_row_idx}: Cột A = '{tot_title}', Cột B = '{tot_formula}'")
    assert tot_title == "TỔNG KẾT", "Dòng tổng kết sai tiêu đề!"
    assert "=AVERAGE" in str(tot_formula), f"Dòng tổng kết thiếu công thức AVERAGE: {tot_formula}"

    wb.close()
    app.destroy()
    print("\n=====================================================================")
    print("=== TẤT CẢ CÁC BÀI KIỂM THỬ QUY TẮC GIỚI TÍNH MỚI ĐÃ PASS 100%! ===")
    print("=====================================================================")

if __name__ == '__main__':
    test_new_gender_mapping()
