# -*- coding: utf-8 -*-
"""
Module: excel_manager.py
Chức năng: Quản lý xuất báo cáo và tự động lưu dữ liệu vào file Excel 'Bao_Cao_Tom.xlsx'
- Định dạng Excel đạt tiêu chuẩn thẩm mỹ chuyên nghiệp bậc cao bằng thư viện openpyxl:
  1. THÔNG TIN CHUNG (Dòng 1 đến Dòng 4):
     - Tiêu đề: "BÁO CÁO DỮ LIỆU ĐẾM TÔM THÔNG MINH" (In hoa, cỡ 16, Bold, Navy Blue #1F4E78).
     - Thông tin: Tên Ao Nuôi, Nhân Viên Cân, Ngày Giờ Xuất Báo Cáo. Nền Light Gray #F2F2F2.
  2. THANH TIÊU ĐỀ BẢNG (Header):
     - Dòng 6: Chiều cao 26, nền Navy Blue #1F4E78, chữ trắng, Calibri 11 Bold, căn giữa.
  3. DÒNG DỮ LIỆU (Data Rows):
     - Bắt đầu từ dòng 7: Chiều cao 20, Calibri 10, viền xám mỏng gọn gàng.
     - STT, Giới Tính, Thời Gian căn giữa. Trọng lượng, Số lượng căn phải định dạng chuẩn.
  4. DÒNG TỔNG KẾT (Total Row):
     - Sử dụng công thức Excel (=AVERAGE(...), =SUM(...)). In đậm, gạch chân kép chuẩn kế toán.
  5. TỰ ĐỘNG GIÃN CỘT:
     - Tự động đo độ dài nội dung để giãn cột, chống che chữ và lỗi "###".
"""

import os
import datetime
from typing import List, Dict, Any, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DEFAULT_EXCEL_FILENAME = "Bao_Cao_Tom.xlsx"

def save_shrimp_report(
    file_path: str,
    pond_name: str,
    staff_name: str,
    weight_records: List[Dict[str, Any]],
    batch_records: List[Dict[str, Any]],
    total_shrimp_count: int
) -> Tuple[bool, str, str]:
    """
    Lưu hoặc cập nhật báo cáo Excel gồm 2 sheet dữ liệu theo tiêu chuẩn thẩm mỹ chuyên nghiệp cao cấp.
    Returns: (success: bool, final_path: str, message: str)
    """
    try:
        wb = Workbook()
        # Xóa sheet mặc định ban đầu của Workbook
        default_sheet = wb.active
        wb.remove(default_sheet)

        # ---------------------------------------------------------------------
        # ĐỊNH NGHĨA PHONG CÁCH & MÀU SẮC CHUẨN OPENPYXL
        # ---------------------------------------------------------------------
        navy_color = '1F4E78'
        white_color = 'FFFFFF'
        gray_fill_color = 'F2F2F2'
        thin_gray_border_color = 'D3D3D3'

        title_font = Font(name='Calibri', size=16, bold=True, color=navy_color)
        meta_label_font = Font(name='Calibri', size=10, bold=True, color='2C3E50')
        meta_val_font = Font(name='Calibri', size=10, bold=True, color=navy_color)

        meta_fill = PatternFill(start_color=gray_fill_color, end_color=gray_fill_color, fill_type='solid')

        header_font = Font(name='Calibri', size=11, bold=True, color=white_color)
        header_fill = PatternFill(start_color=navy_color, end_color=navy_color, fill_type='solid')

        data_font = Font(name='Calibri', size=10, color='000000')
        total_font = Font(name='Calibri', size=10, bold=True, color=navy_color)

        align_center = Alignment(horizontal='center', vertical='center')
        align_right = Alignment(horizontal='right', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')

        thin_gray_side = Side(border_style='thin', color=thin_gray_border_color)
        data_border = Border(
            left=thin_gray_side,
            right=thin_gray_side,
            top=thin_gray_side,
            bottom=thin_gray_side
        )

        header_border = Border(
            left=Side(style='thin', color='FFFFFF'),
            right=Side(style='thin', color='FFFFFF'),
            top=Side(style='medium', color=navy_color),
            bottom=Side(style='medium', color=navy_color)
        )

        # Đường gạch chân kép chuẩn kế toán cho dòng tổng kết (Total Row)
        accounting_border = Border(
            top=Side(style='thin', color=navy_color),
            bottom=Side(style='double', color=navy_color),
            left=thin_gray_side,
            right=thin_gray_side
        )

        now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        safe_pond = pond_name if pond_name else "Ao Mặc Định"
        safe_staff = staff_name if staff_name else "Kỹ Thuật Viên"

        # =====================================================================
        # SHEET 1: CÂN TRỌNG LƯỢNG & PHÂN LOẠI GIỚI TÍNH
        # =====================================================================
        ws1 = wb.create_sheet(title="Can_Trong_Luong_Gioi_Tinh")
        ws1.views.sheetView[0].showGridLines = True

        # 1. THÔNG TIN CHUNG (Dòng 1 đến Dòng 4)
        # Dòng 1: Tiêu đề báo cáo phóng to cỡ chữ 16, In hoa, Bold, Navy Blue
        ws1.merge_cells('A1:D1')
        title_cell_1 = ws1['A1']
        title_cell_1.value = "BÁO CÁO DỮ LIỆU ĐẾM TÔM THÔNG MINH"
        title_cell_1.font = title_font
        title_cell_1.alignment = align_center
        ws1.row_dimensions[1].height = 34

        # Dưới đó ghi rõ: Tên Ao Nuôi, Nhân Viên Cân, Ngày Giờ Xuất Báo Cáo.
        # Khu vực này tô nền màu xám nhạt (Light Gray #F2F2F2)
        meta_rows_1 = [
            ("Tên Ao Nuôi / Mã Lô:", safe_pond),
            ("Nhân Viên Cân:", safe_staff),
            ("Ngày Giờ Xuất Báo Cáo:", now_str)
        ]

        for idx, (label, val) in enumerate(meta_rows_1, start=2):
            ws1.row_dimensions[idx].height = 20
            cell_lbl = ws1.cell(row=idx, column=1, value=label)
            cell_lbl.font = meta_label_font
            cell_lbl.alignment = align_left

            cell_val = ws1.cell(row=idx, column=2, value=val)
            cell_val.font = meta_val_font
            cell_val.alignment = align_left

            # Tô nền Light Gray cho cả khu vực thông tin dòng 2-4
            for col_c in range(1, 5):
                ws1.cell(row=idx, column=col_c).fill = meta_fill

        # Dòng 5: Dòng trống phân cách
        ws1.row_dimensions[5].height = 12

        # 2. ĐỊNH DẠNG THANH TIÊU ĐỀ BẢNG (Header) - Dòng 6
        headers_1 = ["STT", "Trọng Lượng (g)", "Giới Tính", "Thời Gian Ghi Nhận"]
        ws1.row_dimensions[6].height = 26
        for col_idx, h_text in enumerate(headers_1, start=1):
            c = ws1.cell(row=6, column=col_idx, value=h_text)
            c.font = header_font
            c.fill = header_fill
            c.alignment = align_center
            c.border = header_border

        # 3. ĐỊNH DẠNG CÁC DÒNG DỮ LIỆU (Data Rows) - Từ dòng 7
        start_row_1 = 7
        curr_row_1 = start_row_1

        for r in weight_records:
            ws1.row_dimensions[curr_row_1].height = 20

            # Cột 1: STT - Căn giữa
            c1 = ws1.cell(row=curr_row_1, column=1, value=r.get('stt', curr_row_1 - start_row_1 + 1))
            c1.alignment = align_center

            # Cột 2: Trọng lượng - Căn phải, định dạng thập phân 1 chữ số ('0.0')
            w_val = float(r.get('weight', 0.0))
            c2 = ws1.cell(row=curr_row_1, column=2, value=w_val)
            c2.alignment = align_right
            c2.number_format = '0.0'

            # Cột 3: Giới tính - Căn giữa (Male/Female/Chưa rõ)
            g_val = r.get('gender', 'Chưa rõ')
            c3 = ws1.cell(row=curr_row_1, column=3, value=g_val)
            c3.alignment = align_center
            if g_val in ('Male', 'Đực'):
                c3.fill = PatternFill(start_color='EBF1F5', end_color='EBF1F5', fill_type='solid')
            elif g_val in ('Female', 'Cái'):
                c3.fill = PatternFill(start_color='FDE9D9', end_color='FDE9D9', fill_type='solid')

            # Cột 4: Thời gian - Căn giữa
            c4 = ws1.cell(row=curr_row_1, column=4, value=r.get('time', ''))
            c4.alignment = align_center

            for col_c in range(1, 5):
                cell = ws1.cell(row=curr_row_1, column=col_c)
                cell.font = data_font
                cell.border = data_border

            curr_row_1 += 1

        end_row_1 = curr_row_1 - 1
        num_records_1 = len(weight_records)

        # 4. DÒNG TỔNG KẾT (Total Row ở dưới cùng bảng)
        total_row_1 = curr_row_1
        ws1.row_dimensions[total_row_1].height = 24

        tot_c1 = ws1.cell(row=total_row_1, column=1, value="TỔNG KẾT")
        tot_c1.alignment = align_center

        if num_records_1 > 0:
            tot_c2 = ws1.cell(row=total_row_1, column=2, value=f"=AVERAGE(B{start_row_1}:B{end_row_1})")
            tot_c3 = ws1.cell(row=total_row_1, column=3, value=f"{num_records_1} mẫu")
        else:
            tot_c2 = ws1.cell(row=total_row_1, column=2, value=0.0)
            tot_c3 = ws1.cell(row=total_row_1, column=3, value="0 mẫu")

        tot_c2.alignment = align_right
        tot_c2.number_format = '0.0'
        tot_c3.alignment = align_center

        tot_c4 = ws1.cell(row=total_row_1, column=4, value="")

        for col_c in range(1, 5):
            cell = ws1.cell(row=total_row_1, column=col_c)
            cell.font = total_font
            cell.border = accounting_border
            cell.fill = PatternFill(start_color='FAFAFA', end_color='FAFAFA', fill_type='solid')

        # 5. TỰ ĐỘNG GIÃN CỘT CHO SHEET 1
        for col in ws1.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Bỏ qua ô A1 vì nó merge ngang từ A1 đến D1
                if cell.row == 1:
                    continue
                if cell.value is not None:
                    val_str = str(cell.value)
                    if val_str.startswith('='):
                        val_str = "123,456.0"
                    max_len = max(max_len, len(val_str))
            ws1.column_dimensions[col_letter].width = max(max_len + 5, 18)

        # =====================================================================
        # SHEET 2: ĐẾM SỐ LƯỢNG CỘNG DỒN TỐC ĐỘ CAO
        # =====================================================================
        ws2 = wb.create_sheet(title="Dem_So_Luong_Cong_Don")
        ws2.views.sheetView[0].showGridLines = True

        # 1. THÔNG TIN CHUNG (Dòng 1 đến Dòng 4)
        ws2.merge_cells('A1:E1')
        title_cell_2 = ws2['A1']
        title_cell_2.value = "BÁO CÁO DỮ LIỆU ĐẾM TÔM THÔNG MINH"
        title_cell_2.font = title_font
        title_cell_2.alignment = align_center
        ws2.row_dimensions[1].height = 34

        for idx, (label, val) in enumerate(meta_rows_1, start=2):
            ws2.row_dimensions[idx].height = 20
            cell_lbl = ws2.cell(row=idx, column=1, value=label)
            cell_lbl.font = meta_label_font
            cell_lbl.alignment = align_left

            cell_val = ws2.cell(row=idx, column=2, value=val)
            cell_val.font = meta_val_font
            cell_val.alignment = align_left

            for col_c in range(1, 6):
                ws2.cell(row=idx, column=col_c).fill = meta_fill

        ws2.row_dimensions[5].height = 12

        # 2. THANH TIÊU ĐỀ BẢNG (Header) - Dòng 6
        headers_2 = ["STT", "Mã Mẻ", "Số Lượng (Con)", "Tổng Tích Lũy (Con)", "Thời Gian Ghi Nhận"]
        ws2.row_dimensions[6].height = 26
        for col_idx, h_text in enumerate(headers_2, start=1):
            c = ws2.cell(row=6, column=col_idx, value=h_text)
            c.font = header_font
            c.fill = header_fill
            c.alignment = align_center
            c.border = header_border

        # 3. DÒNG DỮ LIỆU (Data Rows) - Từ dòng 7
        start_row_2 = 7
        curr_row_2 = start_row_2

        for r in batch_records:
            ws2.row_dimensions[curr_row_2].height = 20

            # Cột 1: STT - Căn giữa
            c1 = ws2.cell(row=curr_row_2, column=1, value=r.get('stt', curr_row_2 - start_row_2 + 1))
            c1.alignment = align_center

            # Cột 2: Mã Mẻ - Căn giữa
            b_name = r.get('batch_name', f"Mẻ {curr_row_2 - start_row_2 + 1}")
            c2 = ws2.cell(row=curr_row_2, column=2, value=b_name)
            c2.alignment = align_center

            # Cột 3: Số Lượng - Căn phải, định dạng số '#,##0'
            cnt_val = int(r.get('count', 0))
            c3 = ws2.cell(row=curr_row_2, column=3, value=cnt_val)
            c3.alignment = align_right
            c3.number_format = '#,##0'

            # Cột 4: Tổng Tích Lũy - Căn phải, định dạng số '#,##0'
            cum_val = int(r.get('cumulative', cnt_val))
            c4 = ws2.cell(row=curr_row_2, column=4, value=cum_val)
            c4.alignment = align_right
            c4.number_format = '#,##0'

            # Cột 5: Thời Gian - Căn giữa
            c5 = ws2.cell(row=curr_row_2, column=5, value=r.get('time', ''))
            c5.alignment = align_center

            for col_c in range(1, 6):
                cell = ws2.cell(row=curr_row_2, column=col_c)
                cell.font = data_font
                cell.border = data_border
                if r.get('action') == 'UNDO':
                    cell.fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, color='888888', strike=True)

            curr_row_2 += 1

        end_row_2 = curr_row_2 - 1
        num_records_2 = len(batch_records)

        # 4. DÒNG TỔNG KẾT (Total Row ở dưới cùng bảng)
        total_row_2 = curr_row_2
        ws2.row_dimensions[total_row_2].height = 24

        tot2_c1 = ws2.cell(row=total_row_2, column=1, value="")
        tot2_c2 = ws2.cell(row=total_row_2, column=2, value="TỔNG CỘNG")
        tot2_c2.alignment = align_center

        if num_records_2 > 0:
            tot2_c3 = ws2.cell(row=total_row_2, column=3, value=f"=SUM(C{start_row_2}:C{end_row_2})")
            tot2_c4 = ws2.cell(row=total_row_2, column=4, value=total_shrimp_count)
        else:
            tot2_c3 = ws2.cell(row=total_row_2, column=3, value=0)
            tot2_c4 = ws2.cell(row=total_row_2, column=4, value=0)

        tot2_c3.alignment = align_right
        tot2_c3.number_format = '#,##0'

        tot2_c4.alignment = align_right
        tot2_c4.number_format = '#,##0'

        tot2_c5 = ws2.cell(row=total_row_2, column=5, value="")

        for col_c in range(1, 6):
            cell = ws2.cell(row=total_row_2, column=col_c)
            cell.font = total_font
            cell.border = accounting_border
            cell.fill = PatternFill(start_color='FAFAFA', end_color='FAFAFA', fill_type='solid')

        # 5. TỰ ĐỘNG GIÃN CỘT CHO SHEET 2
        for col in ws2.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Bỏ qua ô A1 vì nó merge ngang từ A1 đến E1
                if cell.row == 1:
                    continue
                if cell.value is not None:
                    val_str = str(cell.value)
                    if val_str.startswith('='):
                        val_str = "123,456.0"
                    max_len = max(max_len, len(val_str))
            ws2.column_dimensions[col_letter].width = max(max_len + 5, 18)

        # ---------------------------------------------------------------------
        # LƯU FILE VỚI CƠ CHẾ AN TOÀN CHỐNG PERMISSION ERROR
        # ---------------------------------------------------------------------
        final_path = file_path
        try:
            wb.save(final_path)
            return True, final_path, f"Đã xuất báo cáo Excel thành công tại '{os.path.basename(final_path)}'"
        except PermissionError:
            timestamp_suffix = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            base, ext = os.path.splitext(file_path)
            alt_path = f"{base}_{timestamp_suffix}{ext}"
            wb.save(alt_path)
            return True, alt_path, f"File gốc đang mở trong Excel! Đã tự động lưu bản sao tại '{os.path.basename(alt_path)}'"

    except Exception as e:
        return False, file_path, f"Lỗi khi lưu file Excel: {str(e)}"
