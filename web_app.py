# -*- coding: utf-8 -*-
"""
=============================================================================
ỨNG DỤNG WEB ĐẾM TÔM THÔNG MINH BẰNG GIỌNG NÓI (RESPONSIVE WEB APP)
Tác giả: Antigravity AI & Kỹ sư Thủy sản
Giao diện: Gradio 6.x (Co giãn tối ưu trên mọi màn hình điện thoại & máy tính)
Công nghệ: 
- Micro trình duyệt web di động (Web Audio API / getUserMedia)
- Nhận diện giọng nói Google Speech Recognition (vi-VN)
- Chuyển đổi âm thanh đa định dạng (PyAV: WebM, Opus, Ogg, M4A -> WAV 16kHz)
- Phản hồi giọng nói tiếng Việt tự nhiên (gTTS + PyAV speedup x1.3)
- Xuất file Excel báo cáo thủy sản chuẩn chuyên nghiệp (OpenPyXL)
=============================================================================
"""

import os
import re
import sys
import time
import wave
import hashlib
import tempfile
from datetime import datetime
from typing import Tuple, List, Dict, Any, Optional

import gradio as gr
import speech_recognition as sr
import av
from gtts import gTTS
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Thư mục cache âm thanh
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TTS_CACHE_DIR = os.path.join(BASE_DIR, "tts_cache")
os.makedirs(TTS_CACHE_DIR, exist_ok=True)

# =============================================================================
# BẢNG ÁNH XẠ CHUYỂN ĐỔI CHỮ THÀNH SỐ (VIETNAMESE WORD TO NUMBER)
# =============================================================================
VIETNAMESE_WORD_TO_NUMBER = {
    # 0 - 10
    'không': 0, 'khong': 0, 'zero': 0,
    'một': 1, 'mot': 1, 'mốt': 1,
    'hai': 2,
    'ba': 3,
    'bốn': 4, 'bon': 4, 'tư': 4, 'tu': 4,
    'năm': 5, 'nam': 5, 'lăm': 5, 'lam': 5, 'nhăm': 5, 'nham': 5,
    'sáu': 6, 'sau': 6,
    'bảy': 7, 'bay': 7, 'bẩy': 7, 'bey': 7,
    'tám': 8, 'tam': 8,
    'chín': 9, 'chin': 9,
    'mười': 10, 'muoi': 10, 'chục': 10, 'chuc': 10,

    # 11 - 19
    'mười một': 11, 'mười mot': 11, 'mười mốt': 11,
    'mười hai': 12, 'mười ba': 13,
    'mười bốn': 14, 'mười tư': 14,
    'mười lăm': 15, 'mười năm': 15, 'mười lam': 15,
    'mười sáu': 16,
    'mười bảy': 17, 'mười bẩy': 17,
    'mười tám': 18, 'mười chín': 19,

    # 20 - 29
    'hai mươi': 20, 'hai chục': 20,
    'hai mốt': 21, 'hai một': 21, 'hai hai': 22, 'hai ba': 23,
    'hai tư': 24, 'hai bốn': 24,
    'hai mươi lăm': 25, 'hai mươi năm': 25,
    'hai lăm': 25, 'hai năm': 25, 'hai sáu': 26,
    'hai bảy': 27, 'hai bẩy': 27, 'hai tám': 28, 'hai chín': 29,

    # 30 - 39
    'ba mươi': 30, 'ba chục': 30,
    'ba mốt': 31, 'ba hai': 32, 'ba ba': 33,
    'ba tư': 34, 'ba bốn': 34,
    'ba mươi lăm': 35, 'ba mươi năm': 35,
    'ba lăm': 35, 'ba năm': 35, 'ba sáu': 36,
    'ba bảy': 37, 'ba bẩy': 37, 'ba tám': 38, 'ba chín': 39,

    # 40 - 49
    'bốn mươi': 40, 'bốn chục': 40,
    'bốn mốt': 41, 'bốn hai': 42, 'bốn ba': 43,
    'bốn tư': 44, 'bốn bốn': 44,
    'bốn mươi lăm': 45, 'bốn mươi năm': 45,
    'bốn lăm': 45, 'bốn năm': 45, 'bốn sáu': 46,
    'bốn bảy': 47, 'bốn bẩy': 47, 'bốn tám': 48, 'bốn chín': 49,

    # 50 - 59
    'năm mươi': 50, 'năm chục': 50,
    'năm mốt': 51, 'năm hai': 52, 'năm ba': 53,
    'năm tư': 54, 'năm bốn': 54,
    'năm mươi lăm': 55, 'năm mươi năm': 55,
    'năm lăm': 55, 'năm năm': 55, 'năm sáu': 56,
    'năm bảy': 57, 'năm bẩy': 57, 'năm tám': 58, 'năm chín': 59,

    # 60 - 100
    'sáu mươi': 60, 'sáu chục': 60, 'sáu mươi lăm': 65, 'sáu lăm': 65,
    'bảy mươi': 70, 'bảy chục': 70, 'bảy mươi lăm': 75, 'bảy lăm': 75,
    'tám mươi': 80, 'tám chục': 80, 'tám mươi lăm': 85, 'tám lăm': 85,
    'chín mươi': 90, 'chín chục': 90, 'chín mươi lăm': 95, 'chín lăm': 95,
    'một trăm': 100, 'trăm': 100
}

# =============================================================================
# CHUYỂN ĐỔI ÂM THANH TRÌNH DUYỆT WEB SANG WAV 16KHZ MONO
# =============================================================================
def convert_browser_audio_to_wav(input_path: str) -> Optional[str]:
    """
    Trình duyệt web di động (Android Chrome, iOS Safari) thường gửi file WebM (Opus)
    hoặc M4A/Ogg. Hàm này sử dụng PyAV để giải mã, chuẩn hóa âm lượng và resample thành WAV chuẩn 16kHz mono.
    """
    if not input_path or not os.path.exists(input_path):
        return None
    try:
        import numpy as np
        container = av.open(input_path)
        audio_stream = next((s for s in container.streams if s.type == 'audio'), None)
        if audio_stream is None:
            container.close()
            return None

        fd, out_wav = tempfile.mkstemp(suffix=".wav")
        os.close(fd)

        resampler = av.AudioResampler(format='s16', layout='mono', rate=16000)

        raw_frames = []
        for frame in container.decode(audio_stream):
            resampled_frames = resampler.resample(frame)
            for r_frame in resampled_frames:
                raw_frames.append(r_frame.to_ndarray())

        # Flush resampler
        flushed = resampler.resample(None)
        if flushed:
            for r_frame in flushed:
                raw_frames.append(r_frame.to_ndarray())

        container.close()

        if not raw_frames:
            return None

        # Kết hợp các frame âm thanh thành mảng 1D
        audio_data = np.concatenate(raw_frames, axis=1)[0]

        # Tự động khuếch đại âm thanh nếu micro điện thoại thu quá nhỏ
        max_amp = np.max(np.abs(audio_data))
        if max_amp > 80:  # Không phải khoảng lặng tuyệt đối
            scale = min(4.0, 26000.0 / max_amp)
            if scale > 1.2:
                audio_data = np.clip(audio_data * scale, -32768, 32767).astype(np.int16)

        with wave.open(out_wav, 'wb') as wav_out:
            wav_out.setnchannels(1)
            wav_out.setsampwidth(2)
            wav_out.setframerate(16000)
            wav_out.writeframes(audio_data.tobytes())

        return out_wav
    except Exception as e:
        print(f"[Audio Convert Error]: {e}")
        return None

def transcribe_audio_file(audio_path: str) -> Tuple[str, str]:
    """
    Nhận diện giọng nói tiếng Việt từ file âm thanh người dùng vừa ghi âm trên web.
    Trả về: (văn_bản_nhận_diện, thông_báo_trạng_thái)
    """
    if not audio_path or not os.path.exists(audio_path):
        return "", "Không tìm thấy dữ liệu âm thanh."

    wav_path = convert_browser_audio_to_wav(audio_path)
    if not wav_path or not os.path.exists(wav_path):
        if audio_path.lower().endswith(".wav"):
            wav_path = audio_path
        else:
            return "", "Lỗi xử lý định dạng âm thanh từ trình duyệt."

    recognizer = sr.Recognizer()
    recognizer.dynamic_energy_threshold = True

    try:
        with sr.AudioFile(wav_path) as source:
            # Lọc nhiễu nền môi trường nhẹ
            recognizer.adjust_for_ambient_noise(source, duration=0.2)
            audio_data = recognizer.record(source)

        text = recognizer.recognize_google(audio_data, language="vi-VN")
        return text.strip(), "Nhận diện thành công."
    except sr.UnknownValueError:
        return "", "Không nhận diện được giọng nói (Hãy nói to, rõ và bấm nút Stop ngay sau khi nói xong)."
    except sr.RequestError as e:
        return "", f"Lỗi kết nối mạng máy chủ Google Speech: {e}"
    except Exception as e:
        return "", f"Lỗi nhận diện âm thanh: {e}"
    finally:
        if wav_path and wav_path != audio_path and os.path.exists(wav_path):
            try:
                os.remove(wav_path)
            except Exception:
                pass

# =============================================================================
# BỘ TỔNG HỢP GIỌNG NÓI PHẢN HỒI (TTS TIẾNG VIỆT TĂNG TỐC ĐỘ 1.3X)
# =============================================================================
def generate_tts_feedback_audio(text: str, speed_rate: float = 1.3) -> Optional[str]:
    """
    Tạo file âm thanh đọc lại số liệu tiếng Việt gửi về điện thoại của người dùng:
    - Sử dụng gTTS kết hợp PyAV libavfilter atempo để tăng tốc độ giữ nguyên cao độ tự nhiên.
    - Lưu cache để phản hồi tức thì nếu câu nói lặp lại.
    """
    if not text:
        return None

    try:
        txt_hash = hashlib.md5(text.strip().encode('utf-8')).hexdigest()
        speed_tag = f"{round(speed_rate, 2):.2f}".rstrip('0').rstrip('.')
        cached_wav = os.path.join(TTS_CACHE_DIR, f"{txt_hash}_x{speed_tag}.wav")
        cached_mp3 = os.path.join(TTS_CACHE_DIR, f"{txt_hash}.mp3")

        # 1. Trả về cache có sẵn
        if os.path.exists(cached_wav):
            return cached_wav

        # 2. Tải MP3 gốc từ gTTS nếu chưa có
        if not os.path.exists(cached_mp3):
            tts = gTTS(text=text, lang='vi', slow=False)
            tts.save(cached_mp3)

        if speed_rate == 1.0:
            return cached_mp3

        # 3. Tăng tốc độ bằng PyAV
        from av.filter import Graph
        import numpy as np

        container = av.open(cached_mp3)
        audio_stream = next(s for s in container.streams if s.type == 'audio')
        sample_rate = audio_stream.rate or 24000

        graph = Graph()
        src = graph.add_abuffer(template=audio_stream)
        atempo_filter = graph.add("atempo", str(speed_rate))
        sink = graph.add("abuffersink")
        src.link_to(atempo_filter)
        atempo_filter.link_to(sink)
        graph.configure()

        out_frames = []
        for frame in container.decode(audio_stream):
            graph.push(frame)
            while True:
                try:
                    out_f = graph.pull()
                    out_frames.append(out_f.to_ndarray())
                except Exception:
                    break

        graph.push(None)
        while True:
            try:
                out_f = graph.pull()
                out_frames.append(out_f.to_ndarray())
            except Exception:
                break
        container.close()

        if out_frames:
            data_tempo = np.concatenate(out_frames, axis=1)
            int16_data = data_tempo[0] if data_tempo.dtype == np.int16 else np.clip(data_tempo[0] * 32767, -32768, 32767).astype(np.int16)
            with wave.open(cached_wav, "wb") as wf:
                wf.setnchannels(1)
                wf.setsampwidth(2)
                wf.setframerate(sample_rate)
                wf.writeframes(int16_data.tobytes())
            return cached_wav

        return cached_mp3
    except Exception as e:
        print(f"[TTS Generate Error]: {e}")
        return None

# =============================================================================
# LOGIC PHÂN LOẠI GIỌNG NÓI & BÓC TÁCH DỮ LIỆU
# =============================================================================
def detect_shrimp_gender(text: str) -> str:
    """
    Bẫy từ khóa giới tính thông minh:
    - Đực: "trai", "chai", "đực", "tai", "male", "mail", "meo", "mê", "mel", "trống" -> "Male"
    - Cái: "cái", "gái", "ái", "female", "phi meo", "pheme", "femail", "phi", "mái" -> "Female"
    - Còn lại: "Chưa rõ"
    """
    if not text:
        return "Chưa rõ"
    t = re.sub(r'(\d+)([a-zA-Zà-ỹÀ-Ỹ]+)', r'\1 \2', text.lower().strip())
    t = re.sub(r'([a-zA-Zà-ỹÀ-Ỹ]+)(\d+)', r'\1 \2', t)

    if re.search(r'\b(cái|cai|gái|gai|ái|female|phi\s*meo|phimeo|pheme|femail|phi|mái|mai)\b', t):
        return "Female"
    if re.search(r'\b(trai|chai|đực|duc|tai|male|mail|meo|mê|mel|trống|trong)\b', t):
        return "Male"
    return "Chưa rõ"

def parse_weight_value(text: str) -> Optional[float]:
    """Bóc tách trọng lượng tôm (gram) từ giọng nói."""
    if not text:
        return None
    t = text.lower().strip()

    # Dạng số rõ ràng: "45", "45g", "45 gram", "42.5"
    m = re.search(r'(\d+(?:[.,]\d+)?)\s*(?:gram|g|gam)?', t)
    if m:
        try:
            val_str = m.group(1).replace(',', '.')
            val = float(val_str)
            if 0 < val < 5000:
                return round(val, 2)
        except ValueError:
            pass

    # Bóc tách số bằng chữ tiếng Việt (VD: "ba mươi lăm gram", "bốn lăm", "năm mươi")
    cleaned = re.sub(r'\b(gram|gam|g|con|tôm|mẫu|mau|cân|can)\b', ' ', t)
    cleaned = re.sub(r'[^\w\s]', ' ', cleaned).strip()

    words = cleaned.split()
    for phrase_len in range(min(5, len(words)), 0, -1):
        for i in range(len(words) - phrase_len + 1):
            chunk = " ".join(words[i:i+phrase_len])
            if chunk in VIETNAMESE_WORD_TO_NUMBER:
                v = VIETNAMESE_WORD_TO_NUMBER[chunk]
                if 0 < v < 5000:
                    return float(v)

    # Phân tích số ghép tổng quát
    total = 0
    curr = 0
    for w in words:
        if w in VIETNAMESE_WORD_TO_NUMBER:
            n = VIETNAMESE_WORD_TO_NUMBER[w]
            if n == 100:
                total += (curr if curr > 0 else 1) * 100
                curr = 0
            elif n == 10 and curr > 0:
                curr = curr * 10
            elif curr > 0 and curr % 10 == 0:
                curr += n
            elif curr == 0:
                curr = n
            else:
                curr = curr * 10 + n
        elif w.isdigit():
            curr = int(w)
    total += curr
    if 0 < total < 5000:
        return float(total)

    return None

def parse_accumulate_batch(text: str) -> Tuple[Optional[str], Optional[int]]:
    """
    Phân tích câu lệnh Chế độ 2:
    - ('RESET', None): Xóa sạch toàn bộ
    - ('UNDO', None): Xóa mẻ cuối / Hoàn tác
    - ('ADD', số_con): Thêm mẻ tôm
    """
    if not text:
        return None, None
    t = text.lower().strip()

    # 1. Lệnh Xóa sạch về 0
    if any(k in t for k in ("xóa hết", "xoa het", "xóa tất cả", "xoa tat ca", "reset")):
        return 'RESET', None

    # 2. Lệnh Hoàn tác mẻ cuối
    if any(k in t for k in ("xóa dòng", "xoa dong", "xóa", "xoa", "hủy", "huy", "undo", "bỏ", "bo")):
        return 'UNDO', None

    # 3. Chuỗi số trực tiếp
    cleaned = re.sub(r'\b(mẻ|me|thêm|them|cộng|cong|con|tôm|tom|bốc|boc|đợt|dot)\b', ' ', t)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()

    if cleaned in VIETNAMESE_WORD_TO_NUMBER:
        val = VIETNAMESE_WORD_TO_NUMBER[cleaned]
        if val > 0:
            return 'ADD', val

    m = re.search(r'\b(\d+)\b', cleaned)
    if m:
        try:
            v = int(m.group(1))
            if v > 0:
                return 'ADD', v
        except ValueError:
            pass

    # 4. Số ghép (VD: "hai mươi lăm")
    words = cleaned.split()
    total = 0
    curr = 0
    for w in words:
        if w in VIETNAMESE_WORD_TO_NUMBER:
            n = VIETNAMESE_WORD_TO_NUMBER[w]
            if n == 100:
                total += (curr if curr > 0 else 1) * 100
                curr = 0
            elif n == 10 and curr > 0:
                curr = curr * 10
            elif curr > 0 and curr % 10 == 0:
                curr += n
            elif curr == 0:
                curr = n
            else:
                curr = curr * 10 + n
        elif w.isdigit():
            curr = int(w)
    total += curr
    if total > 0:
        return 'ADD', total

    return None, None

# =============================================================================
# HÀM XUẤT BÁO CÁO EXCEL CHUẨN THỦY SẢN CHUYÊN NGHIỆP
# =============================================================================
def export_excel_report(pond_name: str, staff_name: str, mode1_data: list, mode2_data: list) -> str:
    """Tạo file Excel Bao_Cao_Tom.xlsx theo chuẩn OpenPyXL thẩm mỹ cao."""
    wb = openpyxl.Workbook()
    
    # ------------------ SHEET 1: CÂN MẪU ------------------
    ws1 = wb.active
    ws1.title = "Can_Trong_Luong"
    ws1.views.sheetView[0].showGridLines = True

    # Banner Tiêu đề
    ws1.merge_cells('A1:D1')
    c1 = ws1['A1']
    c1.value = "BÁO CÁO DỮ LIỆU ĐẾM TÔM THÔNG MINH"
    c1.font = Font(name='Calibri', size=16, bold=True, color='1F4E78')
    c1.alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[1].height = 36

    # Thông tin chung
    info_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    info_font = Font(name='Calibri', size=11, bold=True)
    val_font = Font(name='Calibri', size=11)

    meta_rows = [
        ("Tên Ao Nuôi:", pond_name or "Ao Nuôi Số 1"),
        ("Kỹ Thuật Viên:", staff_name or "Nguyễn Văn A"),
        ("Ngày Xuất Báo Cáo:", datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
    ]
    for idx, (label, val) in enumerate(meta_rows, start=2):
        ws1[f'A{idx}'].value = label
        ws1[f'A{idx}'].font = info_font
        ws1[f'A{idx}'].fill = info_fill
        ws1[f'B{idx}'].value = val
        ws1[f'B{idx}'].font = val_font
        ws1[f'B{idx}'].fill = info_fill
        ws1.merge_cells(f'B{idx}:D{idx}')
        ws1.row_dimensions[idx].height = 20

    ws1.row_dimensions[5].height = 10

    # Header bảng
    headers1 = ["Số Thứ Tự", "Giới Tính (Sex)", "Trọng Lượng (g)", "Thời Gian"]
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws1.row_dimensions[6].height = 26
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=6, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Dữ liệu
    row_num = 7
    total_weight = 0.0
    for item in mode1_data:
        ws1.row_dimensions[row_num].height = 20
        c_stt = ws1.cell(row=row_num, column=1, value=item.get("stt", row_num - 6))
        c_stt.alignment = Alignment(horizontal='center', vertical='center')
        
        c_sex = ws1.cell(row=row_num, column=2, value=item.get("gender", "Chưa rõ"))
        c_sex.alignment = Alignment(horizontal='center', vertical='center')
        if item.get("gender") == "Male":
            c_sex.font = Font(name='Calibri', size=10, bold=True, color="1F4E78")
        elif item.get("gender") == "Female":
            c_sex.font = Font(name='Calibri', size=10, bold=True, color="C00000")

        w_val = float(item.get("weight", 0.0))
        total_weight += w_val
        c_wt = ws1.cell(row=row_num, column=3, value=w_val)
        c_wt.alignment = Alignment(horizontal='right', vertical='center')
        c_wt.number_format = '#,##0.00'

        c_time = ws1.cell(row=row_num, column=4, value=item.get("time", ""))
        c_time.alignment = Alignment(horizontal='center', vertical='center')

        for col in range(1, 5):
            ws1.cell(row=row_num, column=col).border = thin_border
        row_num += 1

    # Thống kê tổng kết
    if mode1_data:
        ws1.row_dimensions[row_num].height = 24
        ws1.cell(row=row_num, column=1, value="TỔNG CỘNG").font = Font(name='Calibri', size=11, bold=True)
        ws1.cell(row=row_num, column=1).alignment = Alignment(horizontal='center', vertical='center')
        c_tot = ws1.cell(row=row_num, column=3, value=total_weight)
        c_tot.font = Font(name='Calibri', size=11, bold=True, color="1F4E78")
        c_tot.alignment = Alignment(horizontal='right', vertical='center')
        c_tot.number_format = '#,##0.00'

        count_samples = len(mode1_data)
        avg_w = total_weight / count_samples if count_samples > 0 else 0
        size_kg = 1000 / avg_w if avg_w > 0 else 0

        ws1.cell(row=row_num + 1, column=1, value=f"Số lượng mẫu: {count_samples} con").font = Font(name='Calibri', italic=True)
        ws1.cell(row=row_num + 2, column=1, value=f"Trọng lượng trung bình: {avg_w:.2f} g/con | Size quy chuẩn: {size_kg:.1f} con/kg").font = Font(name='Calibri', bold=True, color="008000")

    # ------------------ SHEET 2: ĐẾM CỘNG DỒN ------------------
    ws2 = wb.create_sheet(title="Dem_Cong_Don")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells('A1:D1')
    c2 = ws2['A1']
    c2.value = "BÁO CÁO ĐẾM TÔM CỘNG DỒN"
    c2.font = Font(name='Calibri', size=16, bold=True, color='1F4E78')
    c2.alignment = Alignment(horizontal='center', vertical='center')
    ws2.row_dimensions[1].height = 36

    for idx, (label, val) in enumerate(meta_rows, start=2):
        ws2[f'A{idx}'].value = label
        ws2[f'A{idx}'].font = info_font
        ws2[f'A{idx}'].fill = info_fill
        ws2[f'B{idx}'].value = val
        ws2[f'B{idx}'].font = val_font
        ws2[f'B{idx}'].fill = info_fill
        ws2.merge_cells(f'B{idx}:D{idx}')
        ws2.row_dimensions[idx].height = 20

    headers2 = ["Số Mẻ (STT)", "Số Lượng Mẻ (Con)", "Tổng Tích Lũy (Con)", "Thời Gian"]
    ws2.row_dimensions[6].height = 26
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=6, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    r2 = 7
    for item in mode2_data:
        ws2.row_dimensions[r2].height = 20
        ws2.cell(row=r2, column=1, value=item.get("stt", r2 - 6)).alignment = Alignment(horizontal='center', vertical='center')
        ws2.cell(row=r2, column=2, value=item.get("batch", 0)).alignment = Alignment(horizontal='right', vertical='center')
        ws2.cell(row=r2, column=3, value=item.get("total", 0)).alignment = Alignment(horizontal='right', vertical='center')
        ws2.cell(row=r2, column=4, value=item.get("time", "")).alignment = Alignment(horizontal='center', vertical='center')
        for col in range(1, 5):
            ws2.cell(row=r2, column=col).border = thin_border
        r2 += 1

    # Điều chỉnh độ rộng cột
    for ws in (ws1, ws2):
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    export_path = os.path.join(BASE_DIR, "Bao_Cao_Tom.xlsx")
    wb.save(export_path)
    return export_path

# =============================================================================
# GIAO DIỆN GRADIO RESPONSIVE (TỐI ƯU HÓA HOÀN TOÀN CHO ĐIỆN THOẠI)
# =============================================================================
CUSTOM_CSS = """
/* Responsive Mobile-First Styling */
body, .gradio-container {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    background-color: #f8fafc;
}

/* Header & Banner */
.app-header {
    background: linear-gradient(135deg, #1e3a8a 0%, #0284c7 100%);
    color: white;
    padding: 18px 20px;
    border-radius: 14px;
    margin-bottom: 16px;
    box-shadow: 0 4px 12px rgba(2, 132, 199, 0.25);
}
.app-header h1 {
    font-size: 1.5rem !important;
    font-weight: 800 !important;
    margin: 0 !important;
    color: white !important;
}
.app-header p {
    font-size: 0.9rem !important;
    opacity: 0.9;
    margin-top: 4px !important;
}

/* Thẻ thống kê KPI Dashboard (To rõ trên di động) */
.kpi-container {
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    gap: 10px;
    margin-bottom: 14px;
}
@media (max-width: 650px) {
    .kpi-container {
        grid-template-columns: repeat(2, 1fr) !important;
        gap: 8px !important;
    }
    .kpi-card {
        padding: 8px 10px !important;
    }
    .kpi-value {
        font-size: 1.35rem !important;
    }
    .kpi-label {
        font-size: 0.72rem !important;
    }
}
.kpi-card {
    background: white;
    border-radius: 12px;
    padding: 12px 14px;
    border: 1px solid #e2e8f0;
    box-shadow: 0 2px 6px rgba(0,0,0,0.04);
    text-align: center;
}
.kpi-label {
    font-size: 0.78rem;
    font-weight: 600;
    color: #64748b;
    text-transform: uppercase;
    letter-spacing: 0.5px;
}
.kpi-value {
    font-size: 1.65rem;
    font-weight: 800;
    margin-top: 2px;
}
.kpi-blue { color: #1e40af; }
.kpi-green { color: #059669; }
.kpi-amber { color: #d97706; }
.kpi-indigo { color: #4338ca; }

/* Nút bấm to, dễ chạm bằng ngón tay trên điện thoại */
button.action-btn {
    border-radius: 12px !important;
    font-size: 1.05rem !important;
    font-weight: 700 !important;
    min-height: 48px !important;
}

/* Bảng dữ liệu */
.dataframe {
    font-size: 0.88rem !important;
}
"""

def create_shrimp_app():
    with gr.Blocks(title="Đếm Tôm Bằng Giọng Nói - Smart Shrimp Counter") as demo:
        # State lưu trữ phiên làm việc
        state_mode1 = gr.State([])  # [{"stt": 1, "gender": "Male", "weight": 45.0, "time": "..."}]
        state_mode2 = gr.State([])  # [{"stt": 1, "batch": 10, "total": 10, "time": "..."}]

        # Header giao diện
        gr.HTML("""
        <div class="app-header">
            <div style="display: flex; align-items: center; justify-content: space-between;">
                <div>
                    <h1>🦐 QUẢN LÝ & ĐẾM TÔM THÔNG MINH</h1>
                    <p>Nhận diện giọng nói tiếng Việt rảnh tay qua micro điện thoại & máy tính</p>
                </div>
            </div>
        </div>
        """)

        # Metadata Thông tin ao nuôi
        with gr.Accordion("⚙️ Thông Tin Ao Nuôi & Kỹ Thuật Viên", open=False):
            with gr.Row():
                txt_pond_name = gr.Textbox(value="Ao Nuôi Số 1", label="Tên Ao Nuôi", scale=1)
                txt_staff_name = gr.Textbox(value="Nguyễn Văn A", label="Kỹ Thuật Viên", scale=1)

        # Tab 2 Chế độ chính
        with gr.Tabs():
            # =================================================================
            # TAB 1: CHẾ ĐỘ 1 - CÂN TRỌNG LƯỢNG & GIỚI TÍNH MẪU
            # =================================================================
            with gr.TabItem("⚖️ CHẾ ĐỘ 1: CÂN MẪU & GIỚI TÍNH"):
                # Dashboard KPI Chế độ 1
                html_kpi_m1 = gr.HTML("""
                <div class="kpi-container">
                    <div class="kpi-card">
                        <div class="kpi-label">Số Mẫu Đã Cân</div>
                        <div class="kpi-value kpi-blue">0 con</div>
                    </div>
                    <div class="kpi-card">
                        <div class="kpi-label">Tổng Trọng Lượng</div>
                        <div class="kpi-value kpi-indigo">0.0 g</div>
                    </div>
                    <div class="kpi-card">
                        <div class="kpi-label">Trọng Lượng TB</div>
                        <div class="kpi-value kpi-amber">0.0 g/con</div>
                    </div>
                    <div class="kpi-card">
                        <div class="kpi-label">Size Tôm Ước Tính</div>
                        <div class="kpi-value kpi-green">0 con/kg</div>
                    </div>
                </div>
                """)

                with gr.Row():
                    with gr.Column(scale=1):
                        gr.HTML("""
                        <div style="background: #e0f2fe; border: 1.5px solid #0284c7; border-radius: 10px; padding: 10px 14px; margin-bottom: 8px; color: #0c4a6e; font-size: 0.95rem; line-height: 1.5;">
                            <strong>📢 HƯỚNG DẪN CÂN MẪU:</strong><br/>
                            1️⃣ Bấm biểu tượng <strong>Micro</strong> ➔ Đọc số liệu (VD: <em>"45 gram đực"</em>, <em>"35 cái"</em>, <em>"xóa"</em>).<br/>
                            2️⃣ Nói xong bấm ngay nút <strong>Stop</strong> (nút đỏ) ➔ <strong>Hệ thống sẽ TỰ ĐỘNG nhận diện & đọc loa ngay!</strong>
                        </div>
                        """)
                        audio_input_m1 = gr.Audio(
                            sources=["microphone", "upload"],
                            type="filepath",
                            label="🎤 Bấm Micro để nói (Nói xong bấm nút Stop đỏ)"
                        )
                        btn_record_m1 = gr.Button("🎙️ Xử Lý Giọng Nói (Hoặc bấm Stop ở trên)", variant="primary", elem_classes=["action-btn"])
                        
                        with gr.Row():
                            txt_manual_m1 = gr.Textbox(
                                placeholder="Hoặc gõ/nói câu lệnh nếu môi trường quá ồn...",
                                show_label=False,
                                scale=4
                            )
                            btn_text_m1 = gr.Button("Gửi", variant="secondary", scale=1)

                        status_box_m1 = gr.Markdown("🟢 Sẵn sàng nhận giọng nói...")

                        with gr.Row():
                            btn_undo_m1 = gr.Button("↩️ Xóa Dòng Cuối", variant="stop")
                            btn_clear_m1 = gr.Button("🗑️ Đặt Lại Về 0", variant="secondary")

                        # Component phát lại âm thanh phản hồi cho điện thoại
                        audio_feedback_m1 = gr.Audio(label="🔊 Phản hồi giọng nói AI", autoplay=True, visible=True)

                    with gr.Column(scale=1):
                        df_history_m1 = gr.Dataframe(
                            headers=["STT", "Giới Tính (Sex)", "Trọng Lượng (g)", "Thời Gian"],
                            datatype=["number", "str", "number", "str"],
                            value=[],
                            label="📋 Danh Sách Cân Mẫu Thời Gian Thực",
                            wrap=True
                        )

            # =================================================================
            # TAB 2: CHẾ ĐỘ 2 - ĐẾM SỐ LƯỢNG CỘNG DỒN TỐC ĐỘ CAO
            # =================================================================
            with gr.TabItem("🔢 CHẾ ĐỘ 2: ĐẾM CỘNG DỒN SỐ LƯỢNG"):
                # Dashboard KPI Chế độ 2
                html_kpi_m2 = gr.HTML("""
                <div class="kpi-container">
                    <div class="kpi-card">
                        <div class="kpi-label">Tổng Số Lượng Tôm</div>
                        <div class="kpi-value kpi-green">0 con</div>
                    </div>
                    <div class="kpi-card">
                        <div class="kpi-label">Mẻ Vừa Đếm</div>
                        <div class="kpi-value kpi-blue">+0 con</div>
                    </div>
                    <div class="kpi-card">
                        <div class="kpi-label">Tổng Số Mẻ</div>
                        <div class="kpi-value kpi-indigo">0 mẻ</div>
                    </div>
                    <div class="kpi-card">
                        <div class="kpi-label">Trung Bình / Mẻ</div>
                        <div class="kpi-value kpi-amber">0 con</div>
                    </div>
                </div>
                """)

                with gr.Row():
                    with gr.Column(scale=1):
                        gr.HTML("""
                        <div style="background: #e0f2fe; border: 1.5px solid #0284c7; border-radius: 10px; padding: 10px 14px; margin-bottom: 8px; color: #0c4a6e; font-size: 0.95rem; line-height: 1.5;">
                            <strong>📢 HƯỚNG DẪN ĐẾM CỘNG DỒN:</strong><br/>
                            1️⃣ Bấm biểu tượng <strong>Micro</strong> ➔ Đọc số con (VD: <em>"năm"</em>, <em>"mười lăm"</em>, <em>"ba lăm"</em>, <em>"xóa"</em>).<br/>
                            2️⃣ Nói xong bấm ngay nút <strong>Stop</strong> (nút đỏ) ➔ <strong>Hệ thống sẽ TỰ ĐỘNG cộng dồn & đọc loa ngay!</strong>
                        </div>
                        """)
                        audio_input_m2 = gr.Audio(
                            sources=["microphone", "upload"],
                            type="filepath",
                            label="🎤 Bấm Micro để nói (Nói xong bấm nút Stop đỏ)"
                        )
                        btn_record_m2 = gr.Button("🎙️ Xử Lý Giọng Nói (Hoặc bấm Stop ở trên)", variant="primary", elem_classes=["action-btn"])
                        
                        with gr.Row():
                            txt_manual_m2 = gr.Textbox(
                                placeholder="Hoặc gõ số lượng (VD: '5', '10', 'xóa')...",
                                show_label=False,
                                scale=4
                            )
                            btn_text_m2 = gr.Button("Gửi", variant="secondary", scale=1)

                        status_box_m2 = gr.Markdown("🟢 Sẵn sàng nhận giọng nói...")

                        with gr.Row():
                            btn_undo_m2 = gr.Button("↩️ Hoàn Tác Mẻ Vừa Rồi", variant="stop")
                            btn_clear_m2 = gr.Button("🗑️ Đặt Lại Về 0", variant="secondary")

                        audio_feedback_m2 = gr.Audio(label="🔊 Phản hồi giọng nói AI", autoplay=True, visible=True)

                    with gr.Column(scale=1):
                        df_history_m2 = gr.Dataframe(
                            headers=["STT", "Số Lượng Mẻ (Con)", "Tổng Tích Lũy (Con)", "Thời Gian"],
                            datatype=["number", "number", "number", "str"],
                            value=[],
                            label="📋 Lịch Sử Các Mẻ Đã Đếm",
                            wrap=True
                        )

        # Xuất Excel Báo Cáo
        with gr.Row():
            btn_export = gr.Button("📥 Xuất File Báo Cáo Excel Chuyên Nghiệp (.xlsx)", variant="primary", elem_classes=["action-btn"])
            file_download = gr.File(label="Tải File Báo Cáo Về Máy", visible=True)

        # =====================================================================
        # HÀM CẬP NHẬT GIAO DIỆN & XỬ LÝ SỰ KIỆN
        # =====================================================================
        def update_kpi_m1_html(data: list) -> str:
            count = len(data)
            total_w = sum(float(x.get("weight", 0.0)) for x in data)
            avg_w = total_w / count if count > 0 else 0.0
            size_kg = (1000.0 / avg_w) if avg_w > 0 else 0.0
            return f"""
            <div class="kpi-container">
                <div class="kpi-card">
                    <div class="kpi-label">Số Mẫu Đã Cân</div>
                    <div class="kpi-value kpi-blue">{count} con</div>
                </div>
                <div class="kpi-card">
                    <div class="kpi-label">Tổng Trọng Lượng</div>
                    <div class="kpi-value kpi-indigo">{total_w:.1f} g</div>
                </div>
                <div class="kpi-card">
                    <div class="kpi-label">Trọng Lượng TB</div>
                    <div class="kpi-value kpi-amber">{avg_w:.2f} g/con</div>
                </div>
                <div class="kpi-card">
                    <div class="kpi-label">Size Tôm Ước Tính</div>
                    <div class="kpi-value kpi-green">{size_kg:.1f} con/kg</div>
                </div>
            </div>
            """

        def update_kpi_m2_html(data: list) -> str:
            count = len(data)
            last_batch = data[-1].get("batch", 0) if count > 0 else 0
            total = data[-1].get("total", 0) if count > 0 else 0
            avg_batch = round(total / count, 1) if count > 0 else 0
            return f"""
            <div class="kpi-container">
                <div class="kpi-card">
                    <div class="kpi-label">Tổng Số Lượng Tôm</div>
                    <div class="kpi-value kpi-green">{total:,} con</div>
                </div>
                <div class="kpi-card">
                    <div class="kpi-label">Mẻ Vừa Đếm</div>
                    <div class="kpi-value kpi-blue">+{last_batch} con</div>
                </div>
                <div class="kpi-card">
                    <div class="kpi-label">Tổng Số Mẻ</div>
                    <div class="kpi-value kpi-indigo">{count} mẻ</div>
                </div>
                <div class="kpi-card">
                    <div class="kpi-label">Trung Bình / Mẻ</div>
                    <div class="kpi-value kpi-amber">{avg_batch} con</div>
                </div>
            </div>
            """

        def m1_data_to_df(data: list):
            rows = []
            for item in data:
                rows.append([item.get("stt"), item.get("gender"), item.get("weight"), item.get("time")])
            return rows

        def m2_data_to_df(data: list):
            rows = []
            for item in data:
                rows.append([item.get("stt"), item.get("batch"), item.get("total"), item.get("time")])
            return rows

        # --- Xử lý sự kiện Chế độ 1 ---
        def process_mode1_input(audio_file, text_in, current_data):
            current_data = list(current_data or [])
            spoken_text = ""
            status_msg = ""

            if audio_file:
                spoken_text, status_msg = transcribe_audio_file(audio_file)
            elif text_in:
                spoken_text = text_in.strip()
                status_msg = "Xử lý văn bản nhập tay."

            if not spoken_text:
                return (
                    current_data,
                    m1_data_to_df(current_data),
                    update_kpi_m1_html(current_data),
                    None,
                    f"⚠️ {status_msg or 'Không nhận được câu lệnh nào.'}"
                )

            t_lower = spoken_text.lower().strip()

            # 1. Kiểm tra lệnh xóa sạch
            if any(k in t_lower for k in ("xóa hết", "xoa het", "xóa tất cả", "xoa tat ca", "reset")):
                current_data.clear()
                tts_path = generate_tts_feedback_audio("Đang xóa toàn bộ dữ liệu")
                return (
                    current_data,
                    [],
                    update_kpi_m1_html(current_data),
                    tts_path,
                    "🗑️ Đã xóa toàn bộ dữ liệu cân mẫu về 0."
                )

            # 2. Kiểm tra lệnh xóa dòng cuối
            if any(k in t_lower for k in ("xóa dòng", "xoa dong", "xóa", "xoa", "hủy", "huy", "undo")):
                if current_data:
                    current_data.pop()
                    tts_path = generate_tts_feedback_audio("Đã xóa dòng cuối")
                    msg = "↩️ Đã xóa dòng dữ liệu cân gần nhất."
                else:
                    tts_path = generate_tts_feedback_audio("Không có dữ liệu để xóa")
                    msg = "⚠️ Bảng dữ liệu đang trống, không thể xóa."
                return (
                    current_data,
                    m1_data_to_df(current_data),
                    update_kpi_m1_html(current_data),
                    tts_path,
                    msg
                )

            # 3. Phân tích trọng lượng & giới tính
            weight = parse_weight_value(spoken_text)
            if weight is None:
                tts_path = generate_tts_feedback_audio("Chưa nghe rõ số trọng lượng")
                return (
                    current_data,
                    m1_data_to_df(current_data),
                    update_kpi_m1_html(current_data),
                    tts_path,
                    f"⚠️ Nghe được: '{spoken_text}' nhưng không tìm thấy số gram hợp lệ."
                )

            gender = detect_shrimp_gender(spoken_text)
            new_stt = len(current_data) + 1
            now_str = datetime.now().strftime("%H:%M:%S")

            current_data.append({
                "stt": new_stt,
                "gender": gender,
                "weight": weight,
                "time": now_str
            })

            # Phát âm thanh phản hồi
            tts_text = f"Đã lưu {int(weight) if weight.is_integer() else weight} gram {gender if gender != 'Chưa rõ' else ''}"
            tts_path = generate_tts_feedback_audio(tts_text.strip())

            status = f"✅ Nhận diện: '{spoken_text}' ➔ Lưu: {weight}g | {gender}"
            return (
                current_data,
                m1_data_to_df(current_data),
                update_kpi_m1_html(current_data),
                tts_path,
                status
            )

        # --- Xử lý sự kiện Chế độ 2 ---
        def process_mode2_input(audio_file, text_in, current_data):
            current_data = list(current_data or [])
            spoken_text = ""
            status_msg = ""

            if audio_file:
                spoken_text, status_msg = transcribe_audio_file(audio_file)
            elif text_in:
                spoken_text = text_in.strip()
                status_msg = "Xử lý văn bản nhập tay."

            if not spoken_text:
                return (
                    current_data,
                    m2_data_to_df(current_data),
                    update_kpi_m2_html(current_data),
                    None,
                    f"⚠️ {status_msg or 'Không nhận được câu lệnh nào.'}"
                )

            action, val = parse_accumulate_batch(spoken_text)

            if action == 'RESET':
                current_data.clear()
                tts_path = generate_tts_feedback_audio("Đang xóa toàn bộ dữ liệu")
                return (
                    current_data,
                    [],
                    update_kpi_m2_html(current_data),
                    tts_path,
                    "🗑️ Đã đặt lại toàn bộ số lượng tôm về 0."
                )

            if action == 'UNDO':
                if current_data:
                    removed = current_data.pop()
                    new_tot = current_data[-1].get("total", 0) if current_data else 0
                    tts_path = generate_tts_feedback_audio(f"Đã hoàn tác, còn lại {new_tot}")
                    msg = f"↩️ Đã trừ mẻ {removed.get('batch')} con. Tổng hiện tại: {new_tot} con."
                else:
                    tts_path = generate_tts_feedback_audio("Không có dữ liệu để hoàn tác")
                    msg = "⚠️ Bảng dữ liệu đang trống, không thể hoàn tác."
                return (
                    current_data,
                    m2_data_to_df(current_data),
                    update_kpi_m2_html(current_data),
                    tts_path,
                    msg
                )

            if action == 'ADD' and val is not None:
                current_tot = current_data[-1].get("total", 0) if current_data else 0
                new_tot = current_tot + val
                new_stt = len(current_data) + 1
                now_str = datetime.now().strftime("%H:%M:%S")

                current_data.append({
                    "stt": new_stt,
                    "batch": val,
                    "total": new_tot,
                    "time": now_str
                })

                tts_text = f"Đã thêm {val} con, tổng {new_tot}"
                tts_path = generate_tts_feedback_audio(tts_text)

                status = f"✅ Nhận diện: '{spoken_text}' ➔ +{val} con (Tổng: {new_tot} con)"
                return (
                    current_data,
                    m2_data_to_df(current_data),
                    update_kpi_m2_html(current_data),
                    tts_path,
                    status
                )

            tts_path = generate_tts_feedback_audio("Chưa nghe rõ số lượng")
            return (
                current_data,
                m2_data_to_df(current_data),
                update_kpi_m2_html(current_data),
                tts_path,
                f"⚠️ Nghe được: '{spoken_text}' nhưng không rõ số lượng con."
            )

        # Gán sự kiện Chế độ 1
        audio_input_m1.stop_recording(
            process_mode1_input,
            inputs=[audio_input_m1, txt_manual_m1, state_mode1],
            outputs=[state_mode1, df_history_m1, html_kpi_m1, audio_feedback_m1, status_box_m1]
        )
        btn_record_m1.click(
            process_mode1_input,
            inputs=[audio_input_m1, txt_manual_m1, state_mode1],
            outputs=[state_mode1, df_history_m1, html_kpi_m1, audio_feedback_m1, status_box_m1]
        )
        btn_text_m1.click(
            process_mode1_input,
            inputs=[audio_input_m1, txt_manual_m1, state_mode1],
            outputs=[state_mode1, df_history_m1, html_kpi_m1, audio_feedback_m1, status_box_m1]
        )
        btn_undo_m1.click(
            lambda d: process_mode1_input(None, "xóa", d),
            inputs=[state_mode1],
            outputs=[state_mode1, df_history_m1, html_kpi_m1, audio_feedback_m1, status_box_m1]
        )
        btn_clear_m1.click(
            lambda d: process_mode1_input(None, "reset", d),
            inputs=[state_mode1],
            outputs=[state_mode1, df_history_m1, html_kpi_m1, audio_feedback_m1, status_box_m1]
        )

        # Gán sự kiện Chế độ 2
        audio_input_m2.stop_recording(
            process_mode2_input,
            inputs=[audio_input_m2, txt_manual_m2, state_mode2],
            outputs=[state_mode2, df_history_m2, html_kpi_m2, audio_feedback_m2, status_box_m2]
        )
        btn_record_m2.click(
            process_mode2_input,
            inputs=[audio_input_m2, txt_manual_m2, state_mode2],
            outputs=[state_mode2, df_history_m2, html_kpi_m2, audio_feedback_m2, status_box_m2]
        )
        btn_text_m2.click(
            process_mode2_input,
            inputs=[audio_input_m2, txt_manual_m2, state_mode2],
            outputs=[state_mode2, df_history_m2, html_kpi_m2, audio_feedback_m2, status_box_m2]
        )
        btn_undo_m2.click(
            lambda d: process_mode2_input(None, "xóa", d),
            inputs=[state_mode2],
            outputs=[state_mode2, df_history_m2, html_kpi_m2, audio_feedback_m2, status_box_m2]
        )
        btn_clear_m2.click(
            lambda d: process_mode2_input(None, "reset", d),
            inputs=[state_mode2],
            outputs=[state_mode2, df_history_m2, html_kpi_m2, audio_feedback_m2, status_box_m2]
        )

        # Gán sự kiện Xuất Excel
        def on_export(pond, staff, m1_d, m2_d):
            path = export_excel_report(pond, staff, m1_d or [], m2_d or [])
            return path

        btn_export.click(
            on_export,
            inputs=[txt_pond_name, txt_staff_name, state_mode1, state_mode2],
            outputs=[file_download]
        )

    return demo

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Shrimp Counter Web App")
    default_port = int(os.getenv("PORT", 7860))
    parser.add_argument("--share", action="store_true", default=True, help="Tạo link công khai HTTPS xxxx.gradio.live để mở trên điện thoại ngoài ao (Mặc định: BẬT)")
    parser.add_argument("--no-share", dest="share", action="store_false", help="Chỉ chạy mạng nội bộ, không tạo link công khai")
    parser.add_argument("--port", type=int, default=default_port, help=f"Cổng mạng (mặc định: {default_port})")
    args = parser.parse_args()

    # Tự động nhận diện môi trường Hugging Face Spaces
    is_hf_space = bool(os.getenv("SPACE_ID"))
    should_share = False if is_hf_space else args.share

    demo = create_shrimp_app()
    print("=" * 65)
    if is_hf_space:
        print(f"🦐 ĐANG CHẠY TRÊN HUGGING FACE SPACES: {os.getenv('SPACE_ID')}")
    else:
        print("🦐 ĐANG KHỞI CHẠY SMART SHRIMP COUNTER WEB APP (RESPONSIVE)...")
        print(f"Truy cập máy tính: http://localhost:{args.port}")
        if should_share:
            print("Đang tạo đường dẫn HTTPS công khai để kết nối Micro từ Điện thoại...")
    print("=" * 65)
    demo.launch(server_name="0.0.0.0", server_port=args.port, share=should_share, theme=gr.themes.Soft(), css=CUSTOM_CSS)
